import datetime
import base64
import json
import os
from pathlib import Path
import re
import time
from io import BytesIO

import fitz
import pandas as pd
import requests
import streamlit as st
from docx import Document
from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage
from openpyxl import load_workbook
from PIL import Image


st.set_page_config(page_title="试验报告自动生成", layout="wide")


# ============================
# Word 中图片尺寸（毫米）
# A4 横向页扣除常见页边距后的可用区域约为 233 x 159 mm。
# 图片高度控制在 130 mm 左右，给图片下方预留约三行文字空间。
# ============================
LANDSCAPE_IMG_WIDTH_MM = 233
LANDSCAPE_IMG_HEIGHT_MM = 130

# ============================
# 合成参数（像素）
# ============================
LANDSCAPE_CANVAS_WIDTH_PX = 3600
LANDSCAPE_CANVAS_HEIGHT_PX = 2010
ROW_GAP_PX = 34
COL_GAP_PX = 34
CANVAS_PADDING_PX = 70
BG_COLOR = (255, 255, 255)
DELETE_PARAGRAPH_MARKER = "__DELETE_EMPTY_REPORT_POSITION__"
DEVICE_TABLE_MARKER = "__INSERT_DEVICE_TABLE_HERE__"
DEFAULT_AI_MODEL = os.environ.get("OPENAI_MODEL", "gpt-5.5")
DEFAULT_AI_BASE_URL = os.environ.get("OPENAI_BASE_URL", "https://sub.deuo.top")
DEFAULT_AI_REASONING_EFFORT = os.environ.get("OPENAI_REASONING_EFFORT", "xhigh")
TEMP_HUMIDITY_MAIN_URL = os.environ.get("TEMP_HUMIDITY_MAIN_URL", "http://prodsc.znzz.yhroot.com").rstrip("/")
TEMP_HUMIDITY_DATA_URL = os.environ.get("TEMP_HUMIDITY_DATA_URL", "http://datacollection-platform.znzz.yhroot.com").rstrip("/")
TEMP_HUMIDITY_STATS_PATH = os.environ.get(
    "TEMP_HUMIDITY_STATS_PATH",
    "/api/temperature-humidity-monitoring/stats/",
)
TEMP_HUMIDITY_TIMEOUT_SECONDS = int(os.environ.get("TEMP_HUMIDITY_TIMEOUT_SECONDS", "45"))
TEMP_HUMIDITY_RETRY_COUNT = int(os.environ.get("TEMP_HUMIDITY_RETRY_COUNT", "3"))
TEMP_CONTROL_URL = os.environ.get("TEMP_CONTROL_URL", "http://shiyantask.znzz.yhroot.com").rstrip("/")
TEMP_CONTROL_PATH = os.environ.get("TEMP_CONTROL_PATH", "/api/temperature-control-instructions/")
TEMP_CONTROL_TIMEOUT_SECONDS = int(os.environ.get("TEMP_CONTROL_TIMEOUT_SECONDS", "8"))
DEVICE_LEDGER_PATH = os.environ.get(
    "DEVICE_LEDGER_PATH",
    r"C:\Users\zhaolianzheng\Desktop\副本热试验设备计量台账20260616.xlsx",
)

COMMISSION_SHEET_KEYWORD = "试验委托单"
SUPPLEMENT_SHEET_KEYWORD = "所需要信息"

PERSONNEL_OPTIONS = [
    "刘灿",
    "李俊波",
    "谢炎君",
    "赵洪暄",
    "苏哲",
    "赵连征",
    "范宽宇",
    "梁欣语",
    "黄旭",
    "邹宏宇",
    "丁何益",
    "陈天易",
    "周晨",
    "郭帅帅",
    "姜贝贝",
]

REPORT_FIELD_ORDER = [
    "报告编号",
    "试验名称",
    "产品名称",
    "产品型号",
    "产品代号",
    "产品编号",
    "委托单号",
    "送试单位",
    "单位地址",
    "试验日期",
    "接收日期",
    "试验目的",
    "试验项目",
    "样品总数",
    "产品外观",
    "试验标准",
    "试验依据技术文件",
    "试验地点温度",
    "试验地点湿度",
    "试验准备时间",
    "试验地点",
    "试验过程",
    "撤收时间",
    "真空是否满足要求",
    "试验真空度",
    "温度是否满足要求",
    "试验温度",
    "时间是否满足要求",
    "试验保持时间",
    "降温是否满足要求",
    "试验降温速率",
    "升温是否满足要求",
    "试验升温速率",
    "循环是否满足要求",
    "试验循环次数",
    "试验备注",
    "流程控制软件",
    "流程控制版本",
    "测控软件",
    "测控版本",
    "试验技术负责人",
    "试验参加者",
    "真空度曲线",
    *[f"试验温度曲线{i:02d}" for i in range(1, 21)],
    "产品试验前状态",
    "产品试验后状态",
]

AUTO_FROM_COMMISSION_FIELDS = [
    "报告编号",
    "试验名称",
    "产品名称",
    "产品型号",
    "产品代号",
    "产品编号",
    "委托单号",
    "送试单位",
    "单位地址",
    "试验项目",
    "样品总数",
    "试验标准",
    "试验依据技术文件",
]

DATE_FIELDS = ["接收日期", "试验准备时间", "撤收时间"]
DATE_RANGE_FIELDS = ["试验日期"]

MANUAL_TEXT_FIELDS = [
    "试验目的",
    "试验地点温度",
    "试验地点湿度",
    "试验过程",
    "试验真空度",
    "试验温度",
    "试验保持时间",
    "试验降温速率",
    "试验升温速率",
    "试验循环次数",
]

AI_EXTRACT_FIELDS = [
    "试验目的",
    "试验真空度",
    "试验温度",
    "试验保持时间",
    "试验降温速率",
    "试验升温速率",
    "试验循环次数",
]

AI_EXTRACT_PROMPT = """
你是热真空试验报告助手。请从用户上传的试验大纲截图或图片中提取报告填写字段。

大纲中可能包含多个试验项目。你必须先根据用户提供的“委托单试验项目/试验名称”在大纲中定位最匹配的试验章节或表格，再只从该章节/表格中提取信息。
如果大纲中有多个相似试验，优先选择名称、类型、试验条件与委托单试验项目最一致的一项。
如果无法匹配到委托单对应的试验项目，不要从其它试验项目中提取数据，字段填空字符串，并在“需要人工确认”说明未匹配到对应试验项目。
只提取截图中明确出现、或可由截图内容直接归纳的信息，不要编造。
温度、真空度、速率、时间、循环次数必须保留单位和符号。
“试验保持时间”“试验降温速率”“试验升温速率”必须精简，每个字段不超过20个汉字/字符，保留最重要数值和单位即可。
“试验保持时间”必须优先识别首末循环和中间循环两个保持时间，并统一写成短格式：
首末**h，中间**h
“试验降温速率”必须包含“或依设备最大能力”，写成短格式，例如：≤5℃/min或依设备最大能力。
“试验升温速率”写成短格式，例如：≤5℃/min。
如果只能识别到一个统一保持时间，必须同时写入首末和中间，例如：首末2h10min，中间2h10min，并在“需要人工确认”说明原文未区分首末循环/中间循环。
如果无法确定，字段填空字符串，并在“需要人工确认”中说明原因。
输出必须是 JSON，不要输出 Markdown。

需要提取的字段：
- 试验目的
- 试验真空度
- 试验温度
- 试验保持时间
- 试验降温速率
- 试验升温速率
- 试验循环次数

JSON 格式：
{
  "匹配到的试验项目": "",
  "试验目的": "",
  "试验真空度": "",
  "试验温度": "",
  "试验保持时间": "",
  "试验降温速率": "",
  "试验升温速率": "",
  "试验循环次数": "",
  "需要人工确认": []
}
""".strip()

TEMP_CONTROL_EXTRACT_PROMPT = """
你是热真空/热循环试验报告助手。请基于网站“试验控温说明”的完整结构化 JSON，综合分析并回答报告需要填写的字段。

要求：
1. 优先使用 experimentNumber 精确匹配的记录，不要混入其它试验编号。
2. 必须阅读整个 JSON 后综合判断，不要只机械搬运单个字段。
3. 重点分析 temperatureRange、temperatureStages、temperatureRate、vacuumRequirement、fieldRemarks、specialInstructions、controlStateInfo、statusChangeInstructions、testType 等所有信息。
4. 只基于控温说明 JSON 中出现的信息回答，不要编造；如果某项只能从大纲补充，就留空。
5. 温度、真空度、速率、时间、循环次数必须保留单位和符号。
6. “试验温度”只填写产品控温高低温范围，不填写热沉/箱温/环境温度要求；如果出现“热沉无要求”等内容必须忽略，例如写成：高温100~104℃、低温-94~-90℃。
7. “试验保持时间”必须同时给出首末循环和中间循环，统一写成“首末**h，中间**h”；如果控温说明只有一个统一保持时间，则首末和中间均使用该时长，例如：首末2h10min，中间2h10min。
8. “试验降温速率”“试验升温速率”优先根据 temperatureRate 和备注综合判断，每个字段不超过20个汉字/字符。
9. 如果控温说明里只有一个统一速率字段或统一速率描述，升温和降温必须使用同一个基准值；降温仅在同一基准值后追加“或依设备最大能力”，例如降温：≤5℃/min或依设备最大能力，升温：≤5℃/min。
10. 如果控温说明无真空要求或写“/”，试验真空度填空，并在“需要人工确认”说明。
11. 输出必须是 JSON，不要输出 Markdown。

JSON 格式：
{
  "匹配到的试验项目": "",
  "试验目的": "",
  "试验真空度": "",
  "试验温度": "",
  "试验保持时间": "",
  "试验降温速率": "",
  "试验升温速率": "",
  "试验循环次数": "",
  "需要人工确认": []
}
""".strip()

HUMIDITY_EXTRACT_PROMPT = """
你是试验大厅温湿度截图识别助手。请从用户上传的温湿度极值统计截图中识别一区试验大厅、二区试验大厅的温度和湿度极值。

只提取截图中明确出现的数据，不要编造。
如果截图里有“最高温/最低温/最高湿/最低湿”，请分别识别。
输出必须是 JSON，不要输出 Markdown。

JSON 格式：
{
  "一区试验大厅": {
    "最高温": "",
    "最低温": "",
    "最高湿": "",
    "最低湿": ""
  },
  "二区试验大厅": {
    "最高温": "",
    "最低温": "",
    "最高湿": "",
    "最低湿": ""
  },
  "需要人工确认": []
}
""".strip()

CURVE_IMAGE_ANALYSIS_PROMPT = """
你是热真空试验报告图片分类助手。用户会上传报告用图片，每张图片前都有一个“文件名”。

主要任务：识别每张图片属于哪一类，并把图片的原始文件名填写到对应报告字段。
需要匹配的字段只有这些：
   - 真空度曲线
   - 试验温度曲线01 到 试验温度曲线20
   - 产品试验前状态
   - 产品试验后状态

辅助任务：如果曲线和试验要求都非常清楚，可以顺便初步判断以下字段是否满足：
   - 真空是否满足要求
   - 温度是否满足要求
   - 时间是否满足要求
   - 降温是否满足要求
   - 升温是否满足要求
   - 循环是否满足要求

规则：
- 温度曲线通常有温度、℃、升温、降温、保温、循环、T、温度通道等信息，应匹配到“试验温度曲线xx”。
- “试验曲线01/02/03”如没有真空/压力特征，默认按温度曲线处理，并顺序匹配到“试验温度曲线01、试验温度曲线02...”
- 真空度、真空曲线、压力曲线、真空压力曲线，或坐标/图中文字包含 Pa、kPa、真空、压力，应匹配到“真空度曲线”。
- 产品照片、外观照片、试验前照片、试验后照片，应根据图中文字或文件名匹配到对应的产品状态字段。
- 字段值必须使用用户上传图片的完整文件名，不要改名，不要省略扩展名。
- 每个“试验温度曲线xx”只能对应一张图片；如果同一组曲线有 1.PNG、1.1.PNG 等多张图片，请按顺序分别匹配到连续的“试验温度曲线01、试验温度曲线02...”。
- 真空度曲线、产品试验前状态、产品试验后状态如确有多张图，可以用数组列出多个文件名。
- 如果无法确定，不要强行匹配，在“需要人工确认”说明。
- 是否满足要求不是主要任务；只有曲线和试验要求都足够清楚才填写“满足”或“不满足”，否则留空或填写“不适用”。
- 输出必须是 JSON，不要输出 Markdown。

JSON 格式：
{
  "图片匹配": {
    "真空度曲线": [],
    "试验温度曲线01": [],
    "试验温度曲线02": [],
    "产品试验前状态": [],
    "产品试验后状态": []
  },
  "是否满足要求": {
    "真空是否满足要求": "",
    "温度是否满足要求": "",
    "时间是否满足要求": "",
    "降温是否满足要求": "",
    "升温是否满足要求": "",
    "循环是否满足要求": ""
  },
  "需要人工确认": []
}
""".strip()

NAME_TEXT_FIELDS = [
    "真空度曲线",
    "产品试验前状态",
    "产品试验后状态",
    *[f"试验温度曲线{i:02d}" for i in range(1, 21)],
]

OPTIONAL_IMAGE_FIELDS = [
    "真空度曲线",
    *[f"试验温度曲线{i:02d}" for i in range(1, 21)],
]

SELECT_FIELD_OPTIONS = {
    "产品外观": ["未发现异常，其它（见照片）", "有异常（见照片）"],
    "试验地点": ["一区试验大厅", "二区试验大厅"],
    "真空是否满足要求": ["满足", "不满足", "不适用"],
    "温度是否满足要求": ["满足", "不满足", "不适用"],
    "时间是否满足要求": ["满足", "不满足", "不适用"],
    "降温是否满足要求": ["满足", "不满足", "不适用"],
    "升温是否满足要求": ["满足", "不满足", "不适用"],
    "循环是否满足要求": ["满足", "不满足", "不适用"],
    "试验备注": [
        "配合客户开展产品性能测试，依托客户地面检测设备开展相关工作，实验室未对该数据的真实性、有效性进行验证，不承担相关责任。",
        "客户自行检测，数据结果由客户自行负责。",
    ],
}

DEVICE_VERSION_OPTIONS = {
    "KM3": {
        "流程控制软件": "KM3流程控制软件",
        "流程控制版本": "Touchvew7.5.1.0",
        "测控软件": "KM3测控软件",
        "测控版本": "dlkw2025-V2.2.1.2.0",
    },
    "KM3L": {
        "流程控制软件": "KM3L流程控制软件",
        "流程控制版本": "ForceContro8.1.0.20",
        "测控软件": "KM3L测控软件",
        "测控版本": "dlkw2025-V2.2.1.0.3",
    },
    "KM5": {
        "流程控制软件": "KM5流程控制软件",
        "流程控制版本": "",
        "测控软件": "KM5测控软件",
        "测控版本": "",
    },
    "BZ1": {
        "流程控制软件": "BZ1流程控制软件",
        "流程控制版本": "Touchvew75.5.1.0",
        "测控软件": "BZ1测控软件",
        "测控版本": "dlkw2025-V2.2.1.0.3",
    },
    "BZ2": {
        "流程控制软件": "BZ2流程控制软件",
        "流程控制版本": "ForceContro8.1.0.20",
        "测控软件": "BZ2测控软件",
        "测控版本": "ForceContro8.1.0.20",
    },
    "CY-1000L-1": {
        "流程控制软件": "1m³快速高低温箱-1控制软件",
        "流程控制版本": "",
        "测控软件": "1m³快速高低温箱-1控制软件",
        "测控版本": "",
    },
    "CY-1000L-2": {
        "流程控制软件": "1m³快速高低温箱-2控制软件",
        "流程控制版本": "",
        "测控软件": "1m³快速高低温箱-2控制软件",
        "测控版本": "",
    },
    "CY-1000L-3": {
        "流程控制软件": "1m³快速高低温箱-3控制软件",
        "流程控制版本": "",
        "测控软件": "1m³快速高低温箱-3控制软件",
        "测控版本": "",
    },
}

DEVICE_LEDGER_CACHE = None
DEVICE_VERSION_ALIAS = {
    "1m³快速高低温箱-1": "CY-1000L-1",
    "1m³快速高低温箱-2": "CY-1000L-2",
    "1m³快速高低温箱-3": "CY-1000L-3",
}

IMAGE_EXTENSION_OPTIONS = ["", ".PNG", ".png", ".JPG", ".jpg", ".JPEG", ".jpeg"]


# ============================
# 通用工具
# ============================
def clean_text(value):
    """统一清理 Excel 单元格内容，避免 None、换行和多余空格影响匹配。"""
    if value is None:
        return ""

    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()

    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")

    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\s*\n\s*", " / ", text)
    return re.sub(r"[ \t\f\v]+", " ", text).strip()


def has_value(value):
    if value is None:
        return False
    if isinstance(value, list):
        return any(has_value(item) for item in value)
    return clean_text(value) != ""


def compact_label(value):
    return re.sub(r"\s+", "", clean_text(value)).replace("：", ":")


def to_compact_date(value):
    text = clean_text(value)
    if not text:
        return ""

    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.datetime.strptime(text[:10], fmt).strftime("%Y%m%d")
        except ValueError:
            continue

    match = re.search(r"(\d{4})\D?(\d{1,2})\D?(\d{1,2})", text)
    if match:
        year, month, day = match.groups()
        return f"{int(year):04d}{int(month):02d}{int(day):02d}"
    return text


def normalize_order_number(value):
    text = clean_text(value)
    text = re.sub(r"\s*-\s*", "-", text)
    return re.sub(r"\s+", "", text)


def infer_report_number(order_number):
    order_number = normalize_order_number(order_number)
    if not order_number:
        return ""
    if re.search(r"[A-Za-z]$", order_number):
        return order_number[:-1] + "c"
    return order_number


def normalize_base_report_number(value):
    text = normalize_order_number(value)
    if not text:
        return ""
    if re.search(r"[bcBC]$", text):
        return text[:-1]
    return text


def linked_report_number(base_number):
    base_number = normalize_base_report_number(base_number)
    return f"{base_number}c" if base_number else ""


def linked_order_number(base_number):
    base_number = normalize_base_report_number(base_number)
    return f"{base_number}b" if base_number else ""


def linked_experiment_number(base_number):
    base_number = normalize_base_report_number(base_number)
    match = re.match(r"^(.*?)(\d{4})(\d{3})$", base_number)
    if match:
        return f"{match.group(1)}{match.group(2)}-{match.group(3)}"
    return base_number


def infer_product_code(product_number):
    codes = []
    for part in re.split(r"\s*/\s*", clean_text(product_number)):
        code = re.sub(r"[（(].*?[）)]", "", part).strip()
        if code:
            codes.append(code)
    return "，".join(codes)


def infer_sample_count(product_number):
    text = clean_text(product_number)
    if not text:
        return ""

    total = 0
    for part in re.split(r"\s*/\s*", text):
        match = re.search(r"[（(](.*?)[）)]", part)
        if match:
            items = [item for item in re.split(r"[、,，;；\s]+", match.group(1)) if item]
            total += len(items)
        elif clean_text(part):
            total += 1

    return str(total) if total else ""


def load_device_ledger(path=DEVICE_LEDGER_PATH):
    global DEVICE_LEDGER_CACHE
    if DEVICE_LEDGER_CACHE is not None and DEVICE_LEDGER_CACHE.get("path") == path:
        return DEVICE_LEDGER_CACHE["devices"]

    devices = []
    if not path or not os.path.exists(path):
        DEVICE_LEDGER_CACHE = {"path": path, "devices": devices}
        return devices

    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = {}
    for cell in ws[1]:
        text = clean_text(cell.value)
        if text:
            header[text] = cell.column

    name_col = header.get("设备名称", 1)
    instrument_col = header.get("仪器名称", 2)
    serial_col = header.get("出厂编号", 3)
    cert_col = header.get("证书编号", 4)
    expiry_col = header.get("有效期", 5)

    current_device = ""
    for row in ws.iter_rows(min_row=2, values_only=True):
        device_name = clean_text(row[name_col - 1]) if len(row) >= name_col else ""
        if device_name:
            current_device = device_name

        instrument = clean_text(row[instrument_col - 1]) if len(row) >= instrument_col else ""
        serial = clean_text(row[serial_col - 1]) if len(row) >= serial_col else ""
        cert = clean_text(row[cert_col - 1]) if len(row) >= cert_col else ""
        expiry = clean_text(row[expiry_col - 1]) if len(row) >= expiry_col else ""

        if not any([current_device, instrument, serial, cert, expiry]):
            continue

        devices.append({
            "设备名称": current_device or instrument,
            "仪器名称": instrument,
            "型号/出厂编号": serial,
            "证书编号": cert,
            "有效期": expiry,
        })

    DEVICE_LEDGER_CACHE = {"path": path, "devices": devices}
    return devices


def device_ledger_options():
    names = []
    for item in load_device_ledger():
        name = clean_text(item.get("设备名称"))
        if name and name not in names:
            names.append(name)
    return names


def build_device_table_for_report(device_name):
    device_name = clean_text(device_name)
    rows = [
        [
            clean_text(item.get("仪器名称") or item.get("设备名称")),
            clean_text(item.get("型号/出厂编号")),
            clean_text(item.get("证书编号")),
            clean_text(item.get("有效期")),
        ]
        for item in load_device_ledger()
        if clean_text(item.get("设备名称")) == device_name
    ]
    return [row for row in rows if any(row)]


def build_device_table_for_report_many(device_names):
    seen = set()
    rows = []
    for device_name in device_names or []:
        for row in build_device_table_for_report(device_name):
            key = tuple(row)
            if key in seen:
                continue
            seen.add(key)
            rows.append(row)
    return rows


def build_device_selection_dataframe(device_names):
    rows = []
    for device_name in device_names or []:
        for row in build_device_table_for_report(device_name):
            rows.append({
                "插入": False,
                "设备名称": row[0],
                "出厂编号/版本": row[1],
                "证书编号": row[2],
                "有效期": row[3],
            })
    if not rows:
        return pd.DataFrame(columns=["插入", "设备名称", "出厂编号/版本", "证书编号", "有效期"])
    return pd.DataFrame(rows)


def infer_test_purpose(context):
    product = "".join(
        item
        for item in [context.get("产品型号", ""), context.get("产品名称", "")]
        if has_value(item)
    )
    test_item = clean_text(context.get("试验项目", ""))

    env_map = {
        "真空热试验": "热真空环境",
        "热真空试验": "热真空环境",
        "热平衡试验": "热平衡环境",
        "振动试验": "振动环境",
        "冲击试验": "冲击环境",
        "高温试验": "高温环境",
        "低温试验": "低温环境",
    }
    environment = env_map.get(test_item, test_item)

    if not product or not environment:
        return ""

    return f"检验{product}承受{environment}的能力，验证产品设计的合理性，暴露产品中元器件、材料和制造工艺的缺陷"


def find_sheet(workbook, keyword):
    for sheet_name in workbook.sheetnames:
        if keyword in sheet_name:
            return workbook[sheet_name]
    return None


def find_value_right(ws, label, max_col_offset=12):
    """在同一行从标签右侧寻找第一个非空值，适配委托单里的合并单元格布局。"""
    target = compact_label(label)

    for row in ws.iter_rows():
        for cell in row:
            if compact_label(cell.value) == target:
                for offset in range(1, max_col_offset + 1):
                    value = ws.cell(cell.row, cell.column + offset).value
                    if has_value(value):
                        return clean_text(value)
    return ""


def find_table_value(ws, header, rows_below=1, max_rows=4):
    """读取“型号/名称/产品编号”等表头下方的值。"""
    target = compact_label(header)

    for row in ws.iter_rows():
        for cell in row:
            if compact_label(cell.value) == target:
                for offset in range(rows_below, rows_below + max_rows):
                    value = ws.cell(cell.row + offset, cell.column).value
                    if has_value(value):
                        return clean_text(value)
    return ""


def read_key_value_sheet(ws):
    """兼容旧的“字段在 A 列、值在 B 列以后”的所需信息 sheet。"""
    context = {}
    key_order = []

    if ws is None:
        return context, key_order

    for row in ws.iter_rows(values_only=True):
        key = clean_text(row[0] if row else "")
        if not key:
            continue

        key_order.append(key)
        values = [clean_text(value) for value in row[1:] if has_value(value)]
        if not values:
            context[key] = ""
        elif len(values) == 1:
            context[key] = values[0]
        else:
            context[key] = values

    return context, key_order


def read_csv_context(data_file):
    data_file.seek(0)
    df = pd.read_csv(data_file, header=None)
    context = {}
    key_order = []

    for _, row in df.iterrows():
        key = clean_text(row.iloc[0])
        if not key:
            continue

        key_order.append(key)
        values = [clean_text(value) for value in row.iloc[1:] if has_value(value)]
        if not values:
            context[key] = ""
        elif len(values) == 1:
            context[key] = values[0]
        else:
            context[key] = values

    return context, key_order


def extract_from_commission_sheet(ws):
    """从“试验委托单（请在此页填写）”自动抽取报告模板常用字段。"""
    context = {
        "试验名称": find_value_right(ws, "试验项目（报告）名称"),
        "委托单号": normalize_order_number(find_value_right(ws, "委托单号")),
        "产品型号": find_table_value(ws, "型号"),
        "产品名称": find_table_value(ws, "名称"),
        "产品编号": find_table_value(ws, "产品编号"),
        "产品尺寸": find_table_value(ws, "尺寸（mm）"),
        "技术状态": find_value_right(ws, "技术状态"),
        "生产单位": find_value_right(ws, "生产单位"),
        "送样方式": find_value_right(ws, "送样方式"),
        "贮存要求": find_value_right(ws, "贮存要求"),
        "报告结论形式": find_value_right(ws, "报告结论形式"),
        "报告提交方式": find_value_right(ws, "报告提交方式"),
        "样品性能测试": find_value_right(ws, "样品性能测试"),
        "产品处置方式": find_value_right(ws, "产品处置方式"),
        "送试单位": find_value_right(ws, "单位名称"),
        "单位地址": find_value_right(ws, "单位地址"),
        "联系人": find_value_right(ws, "联系人"),
        "联系电话": find_value_right(ws, "联系电话"),
        "试验项目": find_value_right(ws, "试验类型"),
        "试验标准": find_value_right(ws, "试验依据标准"),
        "试验依据技术文件": find_value_right(ws, "试验依据技术文件"),
        "试验条件": find_value_right(ws, "试验条件"),
        "特殊要求": find_value_right(ws, "特殊要求"),
    }

    start_date = find_value_right(ws, "预估开始日期")
    end_date = find_value_right(ws, "预估完成日期")
    start_compact = to_compact_date(start_date)
    end_compact = to_compact_date(end_date)

    if start_compact and end_compact:
        context["试验日期"] = f"{start_compact}-{end_compact}"
    elif start_compact:
        context["试验日期"] = start_compact

    context["接收日期"] = start_compact
    context["报告编号"] = infer_report_number(context.get("委托单号", ""))
    context["产品代号"] = infer_product_code(context.get("产品编号", ""))
    context["样品总数"] = infer_sample_count(context.get("产品编号", ""))
    context["试验目的"] = infer_test_purpose(context)

    return {key: value for key, value in context.items() if has_value(value)}


def build_context_from_excel(data_file):
    data_file.seek(0)
    workbook = load_workbook(data_file, data_only=True)

    commission_ws = find_sheet(workbook, COMMISSION_SHEET_KEYWORD) or workbook.worksheets[0]
    supplement_ws = find_sheet(workbook, SUPPLEMENT_SHEET_KEYWORD)

    context = extract_from_commission_sheet(commission_ws)
    supplement_context, _ = read_key_value_sheet(supplement_ws)
    key_order = list(REPORT_FIELD_ORDER)

    # 自动抽取结果优先；如果旧委托单仍带有“所需要信息”sheet，只把它作为可选补充值。
    for key, value in supplement_context.items():
        if not has_value(context.get(key)):
            context[key] = value

    for key in key_order:
        context.setdefault(key, "")

    for key in context:
        if key not in key_order:
            key_order.append(key)

    context["生成时间"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    if "生成时间" not in key_order:
        key_order.append("生成时间")

    return context, key_order, commission_ws.title, supplement_ws.title if supplement_ws else ""


def build_initial_context(data_file):
    if data_file.name.lower().endswith(".csv"):
        context, key_order = read_csv_context(data_file)
        return context, key_order, "CSV", ""

    return build_context_from_excel(data_file)


def context_to_editor_rows(context, template_vars, key_order):
    keys = []
    for key in key_order:
        if key not in keys:
            keys.append(key)

    for key in template_vars:
        if key not in keys:
            keys.append(key)

    for key in context:
        if key not in keys:
            keys.append(key)

    return pd.DataFrame(
        [
            {
                "序号": index,
                "字段": key,
                "内容": "、".join(context[key]) if isinstance(context.get(key), list) else clean_text(context.get(key)),
            }
            for index, key in enumerate(keys, start=1)
        ]
    )


def editor_rows_to_context(df):
    context = {}
    for _, row in df.iterrows():
        key = clean_text(row.get("字段"))
        if key:
            context[key] = normalize_field_value(key, row.get("内容"))
    return context


def normalize_field_value(field, value):
    value = clean_text(value)
    if field == "试验温度":
        return normalize_test_temperature(value)
    if field == "试验保持时间":
        return trim_field_text(normalize_hold_time(value), 20)
    if field == "试验降温速率":
        value = normalize_rate_text(value)
        if value and "或依设备最大能力" not in value:
            value = f"{value}或依设备最大能力"
        return trim_field_text(value, 20)
    if field == "试验升温速率":
        return trim_field_text(normalize_rate_text(value), 20)
    return value


def normalize_test_temperature(value):
    text = clean_text(value)
    if not text:
        return ""

    text = text.replace("产品高温", "高温").replace("产品低温", "低温")
    parts = [part.strip() for part in re.split(r"[;；，,\n\r]+", text) if part.strip()]
    kept = []
    for part in parts:
        if re.search(r"热沉|箱温|环境温度", part):
            continue
        if re.search(r"高温|低温", part) and re.search(r"\d", part):
            kept.append(part)

    if kept:
        return "、".join(kept)

    text = re.sub(r"(?:[;；，,、]?\s*)?(?:热沉|箱温|环境温度)[^;；，,、。]*", "", text)
    return re.sub(r"\s+", "", text).strip("、,，;；。")


def trim_field_text(value, max_len):
    value = clean_text(value).replace("，", ",").replace("。", "")
    return value[:max_len]


def normalize_rate_text(value):
    text = clean_text(value)
    if not text:
        return ""
    text = text.replace("不大于", "≤").replace("小于等于", "≤").replace("不超过", "≤")
    text = text.replace("℃ /", "℃/").replace("°C", "℃").replace("min", "min")
    match = re.search(r"(≤|>=|≥|<|>|=)?\s*([0-9]+(?:\.[0-9]+)?)\s*(?:℃|°C)?\s*/?\s*(?:min|分钟|h|小时)", text, re.I)
    if match:
        symbol = match.group(1) or ""
        number = match.group(2)
        unit = "℃/h" if re.search(r"/?\s*(?:h|小时)", match.group(0), re.I) else "℃/min"
        return f"{symbol}{number}{unit}"
    return re.sub(r"\s+", "", text)


def extract_rate_from_text(value, keywords=None):
    text = clean_text(value)
    if not text:
        return ""

    text = text.replace("不大于", "≤").replace("小于等于", "≤").replace("不超过", "≤")
    text = text.replace("℃ /", "℃/").replace("°C", "℃")

    search_texts = []
    if keywords:
        for keyword in keywords:
            for match in re.finditer(rf"{re.escape(keyword)}[^。；;\n\r]{{0,120}}", text, re.I):
                search_texts.append(match.group(0))
    search_texts.append(text)

    rate_pattern = (
        r"(≤|>=|≥|<|>|=)?\s*([0-9]+(?:\.[0-9]+)?)\s*"
        r"(?:(?:℃|°C)\s*/?\s*(?:min|分钟|h|小时)|/\s*(?:min|分钟|h|小时))"
    )
    for candidate in search_texts:
        match = re.search(rate_pattern, candidate, re.I)
        if match:
            symbol = match.group(1) or ""
            number = match.group(2)
            unit = "℃/h" if re.search(r"/?\s*(?:h|小时)", match.group(0), re.I) else "℃/min"
            return f"{symbol}{number}{unit}"
    return ""


def uploaded_file_signature(*files):
    parts = []
    for file in files:
        if file is None:
            parts.append(("none", 0))
            continue
        parts.append((file.name, getattr(file, "size", 0)))
    return tuple(parts)


def ensure_editor_cache(context, template_vars, key_order, data_signature):
    """同一份上传文件只初始化一次表格，避免编辑内容被 Streamlit 重跑覆盖。"""
    if st.session_state.get("data_signature") != data_signature:
        st.session_state["data_signature"] = data_signature
        st.session_state.pop("context_editor", None)
        st.session_state["editor_df"] = context_to_editor_rows(context, template_vars, key_order)

    if "editor_df" not in st.session_state:
        st.session_state["editor_df"] = context_to_editor_rows(context, template_vars, key_order)

    return st.session_state["editor_df"]


def sync_field_widget_value(field, value):
    value = normalize_field_value(field, value)
    for prefix, fields in [("text", MANUAL_TEXT_FIELDS), ("name", NAME_TEXT_FIELDS)]:
        if field in fields:
            st.session_state[f"{prefix}_{field}"] = value


def queue_widget_sync(updates):
    pending = st.session_state.get("pending_widget_sync", {})
    pending.update({field: normalize_field_value(field, value) for field, value in updates.items()})
    st.session_state["pending_widget_sync"] = pending


def apply_pending_widget_sync():
    pending = st.session_state.pop("pending_widget_sync", {})
    for field, value in pending.items():
        update_editor_value(field, value)
        if field in MANUAL_TEXT_FIELDS:
            st.session_state[f"text_{field}"] = value
        elif field in NAME_TEXT_FIELDS:
            st.session_state[f"name_{field}"] = value
        elif field in SELECT_FIELD_OPTIONS:
            st.session_state[f"select_{field}"] = value


def update_editor_value(field, value, sync_widget=False):
    if "editor_df" not in st.session_state:
        return

    value = normalize_field_value(field, value)
    df = st.session_state["editor_df"].copy()
    mask = df["字段"] == field
    if mask.any():
        df.loc[mask, "内容"] = value
    else:
        df.loc[len(df)] = {
            "序号": len(df) + 1,
            "字段": field,
            "内容": value,
        }
    st.session_state["editor_df"] = df
    if sync_widget:
        sync_field_widget_value(field, value)


def rerun_app():
    if hasattr(st, "rerun"):
        st.rerun()
    else:
        st.experimental_rerun()


def split_people(value):
    text = clean_text(value)
    if not text:
        return []
    return [item.strip() for item in re.split(r"[、,，;；\s]+", text) if item.strip()]


def parse_date_for_input(value):
    text = clean_text(value)
    if not text:
        return None

    match = re.search(r"(\d{4})\D?(\d{1,2})\D?(\d{1,2})", text)
    if not match:
        return None

    year, month, day = match.groups()
    try:
        return datetime.date(int(year), int(month), int(day))
    except ValueError:
        return None


def parse_date_range_for_input(value):
    text = clean_text(value)
    if not text:
        return None

    matches = re.findall(r"(\d{4})\D?(\d{1,2})\D?(\d{1,2})", text)
    if not matches:
        return None

    dates = []
    for year, month, day in matches[:2]:
        try:
            dates.append(datetime.date(int(year), int(month), int(day)))
        except ValueError:
            return None

    if len(dates) == 1:
        dates.append(dates[0])
    return tuple(dates)


def format_date_value(value):
    if isinstance(value, tuple):
        if len(value) == 2:
            start, end = value
            if start and end:
                return f"{start.strftime('%Y%m%d')}-{end.strftime('%Y%m%d')}"
        return ""

    if isinstance(value, list):
        return format_date_value(tuple(value))

    if isinstance(value, datetime.datetime):
        return value.strftime("%Y%m%d")

    if isinstance(value, datetime.date):
        return value.strftime("%Y%m%d")

    return clean_text(value)


def current_context_from_editor():
    return editor_rows_to_context(st.session_state.get("editor_df", pd.DataFrame(columns=["字段", "内容"])))


def image_file_to_data_url(file):
    file.seek(0)
    data = file.read()
    file.seek(0)
    suffix = file.name.lower().rsplit(".", 1)[-1] if "." in file.name else "png"
    mime_map = {
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png",
        "webp": "image/webp",
    }
    mime_type = mime_map.get(suffix, "image/png")
    encoded = base64.b64encode(data).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def bytes_to_image_data_url(data, mime_type="image/png"):
    encoded = base64.b64encode(data).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def pdf_to_text_and_images(file, max_text_chars=12000, max_image_pages=6):
    file.seek(0)
    pdf_bytes = file.read()
    file.seek(0)

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    text_parts = []
    for index, page in enumerate(doc, start=1):
        page_text = page.get_text("text").strip()
        if page_text:
            text_parts.append(f"[PDF第{index}页]\n{page_text}")

    text = "\n\n".join(text_parts)[:max_text_chars]
    image_parts = []

    # 扫描版 PDF 往往没有可复制文字，这时把页面转成图片交给视觉模型。
    if len(text.strip()) < 200:
        for page in list(doc)[:max_image_pages]:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            image_parts.append({
                "type": "input_image",
                "image_url": bytes_to_image_data_url(pix.tobytes("png")),
                "detail": "high",
            })

    doc.close()
    return text, image_parts


def file_to_text_snippet(file, max_chars=12000):
    suffix = Path(file.name).suffix.lower()
    file.seek(0)

    if suffix in [".txt", ".md", ".csv"]:
        text = file.read().decode("utf-8", errors="ignore")
        file.seek(0)
        return text[:max_chars]

    if suffix in [".docx"]:
        doc = Document(file)
        file.seek(0)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for tbl in doc.tables:
            for row in tbl.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts)[:max_chars]

    if suffix in [".xlsx", ".xlsm"]:
        wb = load_workbook(file, data_only=True)
        file.seek(0)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet] {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [clean_text(v) for v in row if has_value(v)]
                if vals:
                    parts.append(" | ".join(vals))
        return "\n".join(parts)[:max_chars]

    if suffix == ".pdf":
        text, _ = pdf_to_text_and_images(file, max_text_chars=max_chars)
        return text

    return ""


def build_ai_outline_payload(outline_files, requirement_context=None):
    text_parts = []
    image_parts = []

    for file in outline_files:
        suffix = Path(file.name).suffix.lower()
        if suffix in [".jpg", ".jpeg", ".png", ".webp"]:
            image_parts.append({
                "type": "input_image",
                "image_url": image_file_to_data_url(file),
                "detail": "high",
            })
        elif suffix == ".pdf":
            text, pdf_image_parts = pdf_to_text_and_images(file)
            if text.strip():
                text_parts.append(f"[文件：{file.name}]\n{text}")
            image_parts.extend(pdf_image_parts)
        else:
            snippet = file_to_text_snippet(file)
            if snippet.strip():
                text_parts.append(f"[文件：{file.name}]\n{snippet}")

    content = [{"type": "input_text", "text": AI_EXTRACT_PROMPT}]
    requirement_context = requirement_context or {}
    target_info = {
        "委托单试验项目": clean_text(requirement_context.get("试验项目")),
        "委托单试验名称": clean_text(requirement_context.get("试验名称")),
        "产品名称": clean_text(requirement_context.get("产品名称")),
        "产品型号": clean_text(requirement_context.get("产品型号")),
    }
    content.append({
        "type": "input_text",
        "text": "请先按以下委托单信息检索并定位大纲中的对应试验项目，再提取该项目的信息：\n"
        + json.dumps(target_info, ensure_ascii=False, indent=2),
    })
    if text_parts:
        content.append({
            "type": "input_text",
            "text": "以下是上传的大纲/附件文本内容，请结合它们提取字段：\n\n" + "\n\n".join(text_parts),
        })
    content.extend(image_parts)
    return content


def responses_api_url(base_url=DEFAULT_AI_BASE_URL):
    base = (base_url or "https://api.openai.com").strip().rstrip("/")
    if base.endswith("/responses"):
        return base
    if base.endswith("/v1"):
        return f"{base}/responses"
    return f"{base}/v1/responses"


def build_responses_payload(model, content, max_output_tokens):
    return {
        "model": model,
        "input": [{"role": "user", "content": content}],
        "text": {"format": {"type": "json_object"}},
        "temperature": 0,
        "max_output_tokens": max_output_tokens,
        "store": False,
        "reasoning": {"effort": DEFAULT_AI_REASONING_EFFORT},
    }


def find_experiment_number(context):
    context = context or {}
    candidates = []
    for field in ["试验编号", "试验单号", "委托单号", "报告编号", "试验名称", "试验项目"]:
        value = clean_text(context.get(field))
        if value:
            candidates.append(value)
    for value in context.values():
        value = clean_text(value)
        if value:
            candidates.append(value)
    text = "\n".join(candidates)
    match = re.search(r"SYZX-[A-Z0-9]+-\d{4}-\d{2,4}", text, re.I)
    return match.group(0).upper() if match else ""


def fetch_temperature_control_instruction(experiment_number):
    experiment_number = clean_text(experiment_number)
    if not experiment_number:
        return {}
    try:
        response = requests.get(
            f"{TEMP_CONTROL_URL}{TEMP_CONTROL_PATH}",
            params={"search": experiment_number},
            headers={"Accept": "application/json"},
            timeout=TEMP_CONTROL_TIMEOUT_SECONDS,
        )
    except requests.RequestException as exc:
        raise RuntimeError("控温说明网站连接超时或不可用，请稍后重试；本次将改用上传的大纲附件。") from exc
    data = request_json(response, "控温说明网站查询")
    rows = data.get("results", []) if isinstance(data, dict) else []
    if not isinstance(rows, list):
        return {}
    exact = [
        row for row in rows
        if isinstance(row, dict) and clean_text(row.get("experimentNumber")).upper() == experiment_number.upper()
    ]
    return exact[0] if exact else (rows[0] if rows and isinstance(rows[0], dict) else {})


def trim_large_payload(value, max_string=3000):
    if isinstance(value, dict):
        return {
            key: trim_large_payload(item, max_string=max_string)
            for key, item in value.items()
            if key not in {"image_base64"}
        }
    if isinstance(value, list):
        return [trim_large_payload(item, max_string=max_string) for item in value]
    if isinstance(value, str) and len(value) > max_string:
        return value[:max_string] + "..."
    return value


def extract_temperature_control_info_with_ai(instruction, api_key, requirement_context=None, model=DEFAULT_AI_MODEL, max_output_tokens=1200):
    if not api_key:
        raise ValueError("请先填写 OpenAI API Key，或设置环境变量 OPENAI_API_KEY。")
    if not instruction:
        raise ValueError("未查询到对应试验编号的控温说明。")

    requirement_context = requirement_context or {}
    compact_instruction = trim_large_payload(instruction)
    content = [
        {"type": "input_text", "text": TEMP_CONTROL_EXTRACT_PROMPT},
        {
            "type": "input_text",
            "text": "委托单/报告上下文：\n" + json.dumps({
                "委托单试验项目": clean_text(requirement_context.get("试验项目")),
                "委托单试验名称": clean_text(requirement_context.get("试验名称")),
                "产品名称": clean_text(requirement_context.get("产品名称")),
                "产品型号": clean_text(requirement_context.get("产品型号")),
                "试验编号": find_experiment_number(requirement_context),
            }, ensure_ascii=False, indent=2),
        },
        {
            "type": "input_text",
            "text": "以下为完整控温说明 JSON（仅去除了图片 base64 大块数据），请基于全部信息综合分析：\n"
            + json.dumps(compact_instruction, ensure_ascii=False, indent=2),
        },
    ]
    payload = build_responses_payload(model, content, max_output_tokens)
    response = requests.post(
        responses_api_url(),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=120,
    )
    if response.status_code >= 400:
        if response.status_code == 401:
            raise RuntimeError("控温说明 AI 提取失败：API Key 无效、过期，或当前 Key 不属于这个接口地址。")
        raise RuntimeError(f"控温说明 AI 提取失败：HTTP {response.status_code}，{api_error_summary(response)}")

    result = parse_json_response(response.json())
    if isinstance(result, dict):
        result["试验温度"] = normalize_field_value("试验温度", result.get("试验温度"))
        result["试验保持时间"] = normalize_hold_time(result.get("试验保持时间"))
        for short_field in ["试验保持时间", "试验降温速率", "试验升温速率"]:
            result[short_field] = normalize_field_value(short_field, result.get(short_field))

        instruction_text = json.dumps(trim_large_payload(instruction, max_string=1000), ensure_ascii=False)
        unified_rate = extract_rate_from_text(instruction.get("temperatureRate")) or extract_rate_from_text(
            instruction_text,
            keywords=["temperatureRate", "升降温", "升温", "降温", "速率", "变温"],
        )
        if unified_rate:
            result["试验升温速率"] = unified_rate
            result["试验降温速率"] = f"{unified_rate}或依设备最大能力"
            for rate_field in ["试验降温速率", "试验升温速率"]:
                result[rate_field] = normalize_field_value(rate_field, result.get(rate_field))

    return {
        field: clean_text(result.get(field))
        for field in AI_EXTRACT_FIELDS
    } | {
        "匹配到的试验项目": clean_text(result.get("匹配到的试验项目") or instruction.get("experimentName") or instruction.get("experimentNumber")),
        "需要人工确认": result.get("需要人工确认", []),
        "来源": "控温说明网站",
        "试验编号": clean_text(instruction.get("experimentNumber")),
    }


def merge_ai_results(primary, fallback):
    primary = primary or {}
    fallback = fallback or {}
    merged = dict(primary)
    for field in AI_EXTRACT_FIELDS:
        if not has_value(merged.get(field)) and has_value(fallback.get(field)):
            merged[field] = fallback.get(field)
    primary_confirm = primary.get("需要人工确认", [])
    fallback_confirm = fallback.get("需要人工确认", [])
    if not isinstance(primary_confirm, list):
        primary_confirm = [primary_confirm] if primary_confirm else []
    if not isinstance(fallback_confirm, list):
        fallback_confirm = [fallback_confirm] if fallback_confirm else []
    merged["需要人工确认"] = primary_confirm + [
        item for item in fallback_confirm if item not in primary_confirm
    ]
    primary_source = clean_text(primary.get("来源"))
    fallback_source = clean_text(fallback.get("来源")) or ("上传大纲" if fallback else "")
    if primary_source and fallback:
        merged["来源"] = f"{primary_source}优先，{fallback_source}补充"
    elif fallback and not primary_source:
        merged["来源"] = fallback_source
    return merged


def api_error_summary(response):
    try:
        data = response.json()
    except ValueError:
        text = response.text
    else:
        if isinstance(data, dict):
            error = data.get("error", data)
            if isinstance(error, dict):
                text = error.get("message") or json.dumps(error, ensure_ascii=False)
            else:
                text = str(error)
        else:
            text = str(data)
    text = re.sub(r"sk-[A-Za-z0-9_\-]+", "sk-***", text or "")
    return text[:300]


def normalize_hold_time(value):
    text = clean_text(value)
    if not text:
        return ""

    first_last_match = re.search(r"(?:首末|首尾|第一.*最后|首.*末)[^0-9\-－]*([0-9]+(?:\.[0-9]+)?(?:\s*h|\s*H|\s*小时)?(?:\s*[0-9]+(?:\.[0-9]+)?\s*(?:min|分钟))?)", text)
    middle_match = re.search(r"(?:中间|中部|其余)[^0-9\-－]*([0-9]+(?:\.[0-9]+)?(?:\s*h|\s*H|\s*小时)?(?:\s*[0-9]+(?:\.[0-9]+)?\s*(?:min|分钟))?)", text)
    if first_last_match and middle_match:
        return f"首末{format_hold_duration(first_last_match.group(1))}，中间{format_hold_duration(middle_match.group(1))}"

    numbers = re.findall(r"([0-9]+(?:\.[0-9]+)?\s*(?:h|H|小时)(?:\s*[0-9]+(?:\.[0-9]+)?\s*(?:min|分钟))?)", text)
    if len(numbers) >= 2 and re.search(r"首末|首尾|中间|其余", text):
        return f"首末{format_hold_duration(numbers[0])}，中间{format_hold_duration(numbers[1])}"

    if "首末循环" in text and "中间循环" in text:
        text = re.sub(r"\s+", "", text)
        text = text.replace("小时", "h").replace("H", "h")
        text = text.replace("首末循环", "首末").replace("中间循环", "中间")
        text = text.replace("：", "").replace(":", "")
        return text.replace("，", ",")

    single_match = re.search(r"([0-9]+(?:\.[0-9]+)?\s*(?:h|H|小时)(?:\s*[0-9]+(?:\.[0-9]+)?\s*(?:min|分钟))?)", text)
    if single_match and re.search(r"首末|首尾|中间|其余|循环|保持", text):
        duration = format_hold_duration(single_match.group(1))
        return f"首末{duration}，中间{duration}"

    return text


def format_hold_duration(value):
    text = clean_text(value)
    text = text.replace("小时", "h").replace("H", "h")
    text = text.replace("分钟", "min")
    text = re.sub(r"\s+", "", text)
    if re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", text):
        return f"{text}h"
    return text


def extract_outline_info_with_ai(outline_files, api_key, requirement_context=None, model=DEFAULT_AI_MODEL, max_output_tokens=1200):
    if not api_key:
        raise ValueError("请先填写 OpenAI API Key，或设置环境变量 OPENAI_API_KEY。")
    if not outline_files:
        raise ValueError("请先上传试验大纲文件。")

    content = build_ai_outline_payload(outline_files, requirement_context=requirement_context)
    if len(content) == 2:
        raise ValueError("上传的大纲文件暂时无法读取内容。当前支持 PDF、图片、Word(.docx)、Excel(.xlsx/.xlsm)、TXT/MD/CSV。")

    payload = build_responses_payload(model, content, max_output_tokens)

    response = requests.post(
        responses_api_url(),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=120,
    )
    if response.status_code >= 400:
        if response.status_code == 401:
            raise RuntimeError("AI 提取失败：API Key 无效、过期，或当前 Key 不属于这个接口地址。请检查页面填写的 Key、OPENAI_API_KEY 和 OPENAI_BASE_URL。")
        raise RuntimeError(f"AI 提取失败：HTTP {response.status_code}，{api_error_summary(response)}")

    result = parse_json_response(response.json())
    if isinstance(result, dict):
        result["试验温度"] = normalize_field_value("试验温度", result.get("试验温度"))
        result["试验保持时间"] = normalize_hold_time(result.get("试验保持时间"))
        for short_field in ["试验保持时间", "试验降温速率", "试验升温速率"]:
            result[short_field] = normalize_field_value(short_field, result.get(short_field))

    return {
        field: clean_text(result.get(field))
        for field in AI_EXTRACT_FIELDS
    } | {
        "匹配到的试验项目": clean_text(result.get("匹配到的试验项目")),
        "需要人工确认": result.get("需要人工确认", [])
    }


def parse_json_response(data):
    output_text = data.get("output_text", "")
    if not output_text:
        text_parts = []
        for item in data.get("output", []):
            for part in item.get("content", []):
                if "text" in part:
                    text_parts.append(part["text"])
        output_text = "\n".join(text_parts).strip()

    if not output_text:
        raise RuntimeError("AI 没有返回可解析文本。")

    try:
        return json.loads(output_text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", output_text, re.S)
        if not match:
            raise RuntimeError(f"AI 返回内容不是 JSON：{output_text}")
        return json.loads(match.group(0))


def extract_humidity_info_with_ai(humidity_files, api_key, model=DEFAULT_AI_MODEL, max_output_tokens=800):
    if not api_key:
        raise ValueError("请先填写 OpenAI API Key，或设置环境变量 OPENAI_API_KEY。")
    if not humidity_files:
        raise ValueError("请先上传温湿度截图。")

    content = [{"type": "input_text", "text": HUMIDITY_EXTRACT_PROMPT}]
    for file in humidity_files:
        content.append({
            "type": "input_image",
            "image_url": image_file_to_data_url(file),
            "detail": "high",
        })

    payload = build_responses_payload(model, content, max_output_tokens)
    response = requests.post(
        responses_api_url(),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=120,
    )
    if response.status_code >= 400:
        if response.status_code == 401:
            raise RuntimeError("温湿度识别失败：API Key 无效、过期，或当前 Key 不属于这个接口地址。请检查页面填写的 Key、OPENAI_API_KEY 和 OPENAI_BASE_URL。")
        raise RuntimeError(f"温湿度识别失败：HTTP {response.status_code}，{api_error_summary(response)}")
    return parse_json_response(response.json())


def analyze_report_images_with_ai(image_files, requirement_context, api_key, model=DEFAULT_AI_MODEL, max_output_tokens=1800):
    if not api_key:
        raise ValueError("请先填写 OpenAI API Key，或设置环境变量 OPENAI_API_KEY。")
    if not image_files:
        raise ValueError("请先上传报告图片。")

    requirement_fields = [
        "试验真空度",
        "试验温度",
        "试验保持时间",
        "试验降温速率",
        "试验升温速率",
        "试验循环次数",
        "试验过程",
    ]
    requirements = {
        field: clean_text(requirement_context.get(field))
        for field in requirement_fields
        if has_value(requirement_context.get(field))
    }
    content = [
        {"type": "input_text", "text": CURVE_IMAGE_ANALYSIS_PROMPT},
        {
            "type": "input_text",
            "text": "当前试验要求 JSON：\n" + json.dumps(requirements, ensure_ascii=False, indent=2),
        },
    ]
    for file in image_files:
        content.append({"type": "input_text", "text": f"文件名：{file.name}"})
        content.append({
            "type": "input_image",
            "image_url": image_file_to_data_url(file),
            "detail": "high",
        })

    payload = build_responses_payload(model, content, max_output_tokens)
    response = requests.post(
        responses_api_url(),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=180,
    )
    if response.status_code >= 400:
        if response.status_code == 401:
            raise RuntimeError("图片识别失败：API Key 无效、过期，或当前 Key 不属于这个接口地址。请检查页面填写的 Key、OPENAI_API_KEY 和 OPENAI_BASE_URL。")
        raise RuntimeError(f"图片识别失败：HTTP {response.status_code}，{api_error_summary(response)}")

    return normalize_report_image_analysis(parse_json_response(response.json()), image_files)


def normalize_report_image_analysis(result, image_files):
    valid_names = {file.name for file in image_files}
    image_matches = result.get("图片匹配", {}) if isinstance(result, dict) else {}
    normalized_matches = {field: [] for field in NAME_TEXT_FIELDS}

    for raw_field, raw_names in image_matches.items():
        field = normalize_image_match_field(raw_field)
        if field not in normalized_matches:
            continue
        if isinstance(raw_names, str):
            names = [raw_names]
        elif isinstance(raw_names, list):
            names = raw_names
        else:
            continue
        normalized_matches[field].extend([
            clean_text(name)
            for name in names
            if clean_text(name) in valid_names and clean_text(name) not in normalized_matches[field]
        ])
    normalized_matches = split_temperature_curve_matches(normalized_matches)

    satisfaction = result.get("是否满足要求", {}) if isinstance(result, dict) else {}
    normalized_satisfaction = {}
    for field in SELECT_FIELD_OPTIONS:
        if field not in satisfaction or not field.endswith("是否满足要求"):
            continue
        value = normalize_satisfaction_value(satisfaction.get(field))
        if value:
            normalized_satisfaction[field] = value

    return {
        "图片匹配": normalized_matches,
        "是否满足要求": normalized_satisfaction,
        "需要人工确认": result.get("需要人工确认", []) if isinstance(result, dict) else [],
    }


def split_temperature_curve_matches(matches):
    temperature_fields = [f"试验温度曲线{i:02d}" for i in range(1, 21)]
    flattened = []
    for field in temperature_fields:
        for name in matches.get(field, []):
            if name not in flattened:
                flattened.append(name)
        matches[field] = []

    for index, name in enumerate(flattened[:len(temperature_fields)]):
        matches[temperature_fields[index]] = [name]
    return matches


def normalize_satisfaction_value(value):
    text = clean_text(value)
    if not text:
        return ""
    if "不适用" in text or "无需" in text:
        return "不适用"
    if "不满足" in text or "未满足" in text or "不合格" in text:
        return "不满足"
    if "满足" in text or "合格" in text:
        return "满足"
    return text if text in ["满足", "不满足", "不适用"] else ""


def normalize_image_match_field(field):
    text = clean_text(field)
    if text in NAME_TEXT_FIELDS:
        return text
    if re.search(r"真空|压力", text):
        return "真空度曲线"
    if re.search(r"试验前|试前|前状态", text):
        return "产品试验前状态"
    if re.search(r"试验后|试后|后状态", text):
        return "产品试验后状态"
    match = re.search(r"(?:试验)?(?:温度)?曲线\s*0?([1-9]|1[0-9]|20)", text)
    if match:
        return f"试验温度曲线{int(match.group(1)):02d}"
    if re.search(r"温度|试验曲线|曲线", text):
        return "试验温度曲线01"
    return text


def extract_number_text(value):
    text = clean_text(value)
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return match.group(0) if match else ""


def format_range(low, high, unit):
    low = extract_number_text(low)
    high = extract_number_text(high)
    if low and high:
        return f"{low}~{high}{unit}"
    if low:
        return f"{low}{unit}"
    if high:
        return f"{high}{unit}"
    return ""


def format_humidity_for_hall(result, hall_name):
    hall = result.get(hall_name, {}) if isinstance(result, dict) else {}
    temperature = format_range(hall.get("最低温"), hall.get("最高温"), "℃")
    humidity = format_range(hall.get("最低湿"), hall.get("最高湿"), "%")
    return temperature, humidity


def format_humidity_preview_value(value):
    return value if has_value(value) else "无数据"


def encode_basic_token(username, password):
    raw = f"{username}:{password}".encode("utf-8")
    return base64.b64encode(raw).decode("ascii")


def normalize_basic_token(value):
    text = clean_text(value)
    if not text:
        return ""
    text = re.sub(r"^basic\s+", "", text, flags=re.I).strip()
    if ":" in text:
        return encode_basic_token(*text.split(":", 1))
    return text


def request_json(response, action_name):
    if response.status_code >= 400:
        raise RuntimeError(f"{action_name}失败：HTTP {response.status_code}，{api_error_summary(response)}")
    try:
        return response.json()
    except ValueError as exc:
        raise RuntimeError(f"{action_name}失败：接口未返回 JSON。") from exc


def login_temperature_humidity_site(username, password, org_name=""):
    username = clean_text(username)
    password = "" if password is None else str(password)
    if not username or not password:
        raise ValueError("请填写温湿度网站账号和密码，或直接填写 Basic Token。")

    session = requests.Session()
    login_url = f"{TEMP_HUMIDITY_MAIN_URL}/api/auth/login/"
    payload = {"username": username, "password": password}
    try:
        response = session.post(login_url, json=payload, timeout=12)
    except requests.RequestException as exc:
        raise RuntimeError("温湿度主系统连接超时或不可用，请确认内网/VPN/网站状态；本次已跳过温湿度写入。") from exc
    if response.status_code in (401, 403):
        raise RuntimeError(
            "温湿度网站登录失败：主系统未返回可用于子系统的 basic_token。"
            "请确认账号是登录名（通常不是中文姓名），或在浏览器登录后手动复制 auth.basic 到 Basic Token。"
        )
    data = request_json(response, "温湿度网站登录")
    token = normalize_basic_token(data.get("basic_token") or data.get("token") or data.get("data", {}).get("basic_token"))
    if not token:
        raise RuntimeError("温湿度网站登录成功但未返回 basic_token，请改用 Basic Token 手动填写。")

    resolved_org = clean_text(org_name) or clean_text(data.get("org_name") or data.get("data", {}).get("org_name"))
    if not resolved_org:
        resolved_org = fetch_main_system_org_name(token)
    return token, resolved_org


def fetch_main_system_context(token):
    token = normalize_basic_token(token)
    if not token:
        return {}
    headers = {"Authorization": f"Basic {token}"}
    try:
        response = requests.get(f"{TEMP_HUMIDITY_MAIN_URL}/api/auth/me", headers=headers, timeout=20)
        if response.status_code >= 400:
            return {}
        data = response.json()
        if isinstance(data, dict):
            payload = data.get("data", data)
            if isinstance(payload, dict):
                org_name = clean_text(
                    payload.get("org_name")
                    or payload.get("selected_org_name")
                    or payload.get("orgName")
                    or payload.get("organization")
                )
                return {
                    "orgName": org_name,
                    "orgId": clean_text(
                        payload.get("org_id")
                        or payload.get("orgId")
                        or payload.get("person_org_id")
                        or payload.get("person_dept_id")
                    ),
                    "userName": clean_text(payload.get("name") or payload.get("username") or payload.get("userName")),
                    "userId": clean_text(
                        payload.get("person_id")
                        or payload.get("user_id")
                        or payload.get("userId")
                        or payload.get("id")
                    ),
                }
    except Exception:
        return {}
    return {}


def fetch_main_system_org_name(token):
    return clean_text(fetch_main_system_context(token).get("orgName"))


def build_temperature_humidity_headers(token="", org_name="", person_id="", include_auth=True):
    token = normalize_basic_token(token)
    if include_auth and not token:
        raise ValueError("缺少温湿度网站登录凭据。")
    headers = {
        "Accept": "application/json, text/plain, */*",
        "Origin": TEMP_HUMIDITY_DATA_URL,
        "Referer": f"{TEMP_HUMIDITY_DATA_URL}/equipment-integration-control/temperature-humidity-monitoring",
    }
    if include_auth and token:
        headers["Authorization"] = f"Basic {token}"
    org_name = clean_text(org_name)
    if org_name:
        headers["X-Org-Name"] = requests.utils.quote(org_name, safe="")
    person_id = clean_text(person_id)
    if person_id:
        headers["X-Person-Id"] = person_id
    return headers


def apply_resolved_temperature_humidity_context(resolved, data):
    if not isinstance(data, dict):
        return resolved
    organization = data.get("organization") if isinstance(data.get("organization"), dict) else {}
    personnel = data.get("personnel") if isinstance(data.get("personnel"), dict) else {}
    resolved["orgName"] = clean_text(organization.get("name")) or resolved.get("orgName", "")
    resolved["personId"] = clean_text(personnel.get("id")) or resolved.get("personId", "")
    return resolved


def resolve_temperature_humidity_context(token, org_name=""):
    context = fetch_main_system_context(token)
    if org_name and not context.get("orgName"):
        context["orgName"] = clean_text(org_name)
    if not context.get("orgName"):
        return {"orgName": clean_text(org_name), "personId": ""}

    payload = {
        "orgName": context.get("orgName", ""),
        "orgId": context.get("orgId", ""),
        "userName": context.get("userName", ""),
        "userId": context.get("userId", ""),
    }
    resolved = {"orgName": clean_text(context.get("orgName")), "personId": clean_text(context.get("userId"))}
    url = f"{TEMP_HUMIDITY_DATA_URL}/api/department/resolve-context/"
    for include_auth in (False, True):
        try:
            headers = build_temperature_humidity_headers(token, context.get("orgName"), include_auth=include_auth)
            response = requests.post(url, json=payload, headers=headers, timeout=20)
            if response.status_code < 400:
                resolved = apply_resolved_temperature_humidity_context(resolved, response.json())
                if resolved.get("personId"):
                    break
        except Exception:
            pass
    return resolved


def build_temperature_humidity_params(start_date, end_date, person_id=""):
    params = {"start_date": start_date, "end_date": end_date}
    person_id = clean_text(person_id)
    if person_id:
        params["personId"] = person_id
    return params


def is_number(value):
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def request_temperature_humidity_stats(url, start_date, end_date, token, org_name, person_id, include_auth=True):
    headers = build_temperature_humidity_headers(token, org_name, person_id, include_auth=include_auth)
    params = build_temperature_humidity_params(start_date, end_date, person_id)
    last_response = None
    last_error = None
    for attempt in range(max(1, TEMP_HUMIDITY_RETRY_COUNT)):
        try:
            response = requests.get(
                url,
                params=params,
                headers=headers,
                timeout=TEMP_HUMIDITY_TIMEOUT_SECONDS,
            )
            last_response = response
            if response.status_code not in (502, 503, 504):
                return response
        except requests.RequestException as exc:
            last_error = exc
        if attempt < TEMP_HUMIDITY_RETRY_COUNT - 1:
            time.sleep(1.5 * (attempt + 1))
    if last_response is not None:
        return last_response
    raise RuntimeError(f"温湿度网站取数失败：连接温湿度网站超时，请稍后重试。{last_error}")


def fetch_temperature_humidity_stats(start_date, end_date, token=None, username="", password="", org_name=""):
    token_from_input = has_value(token)
    if not token_from_input:
        token, org_name = login_temperature_humidity_site(username, password, org_name)
    org_name = clean_text(org_name) or fetch_main_system_org_name(token)
    context = resolve_temperature_humidity_context(token, org_name)
    org_name = clean_text(context.get("orgName")) or org_name
    person_id = clean_text(context.get("personId"))
    url = f"{TEMP_HUMIDITY_DATA_URL}{TEMP_HUMIDITY_STATS_PATH}"
    if person_id:
        response = request_temperature_humidity_stats(url, start_date, end_date, token, org_name, person_id, include_auth=False)
    else:
        response = request_temperature_humidity_stats(url, start_date, end_date, token, org_name, person_id)
    if response.status_code in (401, 403) and not person_id:
        response = request_temperature_humidity_stats(url, start_date, end_date, token, org_name, person_id, include_auth=False)
    if response.status_code in (401, 403) and token_from_input and has_value(username) and has_value(password):
        token, org_name = login_temperature_humidity_site(username, password, org_name)
        context = resolve_temperature_humidity_context(token, org_name)
        person_id = clean_text(context.get("personId"))
        if person_id:
            response = request_temperature_humidity_stats(url, start_date, end_date, token, org_name, person_id, include_auth=False)
        else:
            response = request_temperature_humidity_stats(url, start_date, end_date, token, org_name, person_id)
            if response.status_code in (401, 403):
                response = request_temperature_humidity_stats(url, start_date, end_date, token, org_name, person_id, include_auth=False)
    if response.status_code in (401, 403) and not clean_text(org_name):
        resolved_org = fetch_main_system_org_name(token)
        if resolved_org:
            context = resolve_temperature_humidity_context(token, resolved_org)
            person_id = clean_text(context.get("personId"))
            if person_id:
                response = request_temperature_humidity_stats(url, start_date, end_date, token, resolved_org, person_id, include_auth=False)
            else:
                response = request_temperature_humidity_stats(url, start_date, end_date, token, resolved_org, person_id)
                if response.status_code in (401, 403):
                    response = request_temperature_humidity_stats(url, start_date, end_date, token, resolved_org, person_id, include_auth=False)
    if response.status_code in (401, 403) and token_from_input:
        raise RuntimeError(
            "温湿度网站取数失败：Basic Token 无效或已过期。"
            "请清空 Basic Token 后只填网站登录名和密码，或重新从浏览器复制 auth.basic。"
        )
    if response.status_code in (401, 403):
        detail = response.text.strip()
        raise RuntimeError(
            "温湿度网站取数失败：温湿度子系统拒绝当前登录上下文。"
            f"当前组织：{org_name or '未识别'}。"
            "请确认组织为“环境试验中心”、账号填写的是登录名；若仍失败，请在浏览器重新登录后复制 auth.basic 到 Basic Token。"
            f"接口返回：{detail[:200]}"
        )
    if response.status_code in (502, 503, 504):
        raise RuntimeError(
            "温湿度网站取数失败：网站后端连接 thingsboard 数据源超时。"
            "这不是账号密码问题，请稍后再点一次获取；程序已自动重试过。"
            f"接口返回：{api_error_summary(response)}"
        )
    data = request_json(response, "温湿度网站取数")
    return normalize_temperature_humidity_stats(data)


def normalize_temperature_humidity_stats(data):
    stats = data.get("stats", data) if isinstance(data, dict) else {}
    devices = data.get("devices", []) if isinstance(data, dict) else []
    result = {}
    for index, hall_name in enumerate(["一区试验大厅", "二区试验大厅"], start=1):
        device_key = f"device{index}"
        device_stats = stats.get(device_key, {}) if isinstance(stats, dict) else {}
        overall = device_stats.get("overall", {}) if isinstance(device_stats, dict) else {}
        temperature = overall.get("temperature", {}) if isinstance(overall, dict) else {}
        humidity = overall.get("humidity", {}) if isinstance(overall, dict) else {}
        source_note = ""
        if not has_value(temperature.get("max")) or not has_value(temperature.get("min")):
            daily_temperature = device_stats.get("daily", {}).get("temperature", []) if isinstance(device_stats, dict) else []
            daily_values = []
            if isinstance(daily_temperature, list):
                for item in daily_temperature:
                    if isinstance(item, dict):
                        daily_values.extend([item.get("max"), item.get("min")])
            daily_numbers = [float(value) for value in daily_values if is_number(value)]
            if daily_numbers:
                temperature = {"max": max(daily_numbers), "min": min(daily_numbers)}
                source_note = "温度由每日明细补算"
        if not has_value(humidity.get("max")) or not has_value(humidity.get("min")):
            daily_humidity = device_stats.get("daily", {}).get("humidity", []) if isinstance(device_stats, dict) else []
            daily_values = []
            if isinstance(daily_humidity, list):
                for item in daily_humidity:
                    if isinstance(item, dict):
                        daily_values.extend([item.get("max"), item.get("min")])
            daily_numbers = [float(value) for value in daily_values if is_number(value)]
            if daily_numbers:
                humidity = {"max": max(daily_numbers), "min": min(daily_numbers)}
                source_note = "湿度由每日明细补算" if not source_note else f"{source_note}；湿度由每日明细补算"
        if not any(is_number(value) for value in [temperature.get("max"), temperature.get("min"), humidity.get("max"), humidity.get("min")]):
            source_note = "网站该日期范围无数据" if not source_note else f"{source_note}；网站该日期范围无数据"
        alias = ""
        if isinstance(devices, list) and len(devices) >= index and isinstance(devices[index - 1], dict):
            alias = clean_text(devices[index - 1].get("alias"))
        result[hall_name] = {
            "设备名称": alias,
            "最高温": temperature.get("max", ""),
            "最低温": temperature.get("min", ""),
            "最高湿": humidity.get("max", ""),
            "最低湿": humidity.get("min", ""),
            "备注": source_note,
        }
    result["需要人工确认"] = []
    for hall_name in ["一区试验大厅", "二区试验大厅"]:
        hall_data = result.get(hall_name, {})
        if "网站该日期范围无数据" in clean_text(hall_data.get("备注")):
            result["需要人工确认"].append(f"{hall_name} 在所选日期范围内未返回温湿度数据。")
    if not any(has_value(result[hall].get("最高温")) or has_value(result[hall].get("最低温")) for hall in result if hall != "需要人工确认"):
        result["需要人工确认"].append("网站未返回温湿度极值，请检查日期范围、账号权限或接口地址。")
    return result


def default_date_range_from_context(context):
    start, end = parse_date_range_for_input(context.get("试验日期"))
    if start and end:
        return start, end
    today = datetime.date.today()
    return today - datetime.timedelta(days=6), today


def temp_humidity_date_range_from_context(context):
    start = parse_date_for_input(context.get("试验准备时间"))
    end = parse_date_for_input(context.get("撤收时间"))
    if start and end:
        return start, end
    return default_date_range_from_context(context)


# ============================
# 图片处理
# ============================
def make_a4_landscape_image(image_files):
    """多图合成为适配 A4 横向页面的图片，优先横向排布，放不下再换行。"""
    images = [Image.open(img).convert("RGB") for img in image_files]
    if not images:
        return None

    max_w = LANDSCAPE_CANVAS_WIDTH_PX - CANVAS_PADDING_PX * 2
    max_h = LANDSCAPE_CANVAS_HEIGHT_PX - CANVAS_PADDING_PX * 2

    best_layout = None
    max_cols = min(len(images), 6)
    for cols in range(max_cols, 0, -1):
        rows = (len(images) + cols - 1) // cols
        cell_w = (max_w - COL_GAP_PX * (cols - 1)) / cols
        cell_h = (max_h - ROW_GAP_PX * (rows - 1)) / rows
        scale = min(
            min(cell_w / img.width, cell_h / img.height)
            for img in images
        )
        score = scale * (1 + cols * 0.05)
        if best_layout is None or score > best_layout["score"]:
            best_layout = {
                "cols": cols,
                "rows": rows,
                "cell_w": cell_w,
                "cell_h": cell_h,
                "scale": scale,
                "score": score,
            }

    cols = best_layout["cols"]
    rows = best_layout["rows"]
    cell_w = best_layout["cell_w"]
    cell_h = best_layout["cell_h"]
    scale = best_layout["scale"]

    resized = [
        img.resize(
            (max(1, int(img.width * scale)), max(1, int(img.height * scale))),
            Image.LANCZOS,
        )
        for img in images
    ]

    canvas = Image.new("RGB", (LANDSCAPE_CANVAS_WIDTH_PX, LANDSCAPE_CANVAS_HEIGHT_PX), BG_COLOR)
    grid_w = cols * cell_w + (cols - 1) * COL_GAP_PX
    grid_h = rows * cell_h + (rows - 1) * ROW_GAP_PX
    start_x = int((LANDSCAPE_CANVAS_WIDTH_PX - grid_w) / 2)
    start_y = int((LANDSCAPE_CANVAS_HEIGHT_PX - grid_h) / 2)

    for index, img in enumerate(resized):
        row = index // cols
        col = index % cols
        cell_x = start_x + int(col * (cell_w + COL_GAP_PX))
        cell_y = start_y + int(row * (cell_h + ROW_GAP_PX))
        x = cell_x + int((cell_w - img.width) / 2)
        y = cell_y + int((cell_h - img.height) / 2)
        canvas.paste(img, (x, y))

    out = BytesIO()
    canvas.save(out, format="JPEG", quality=95)
    out.seek(0)
    return out


def split_image_names(value, image_map):
    if isinstance(value, list):
        names = [clean_text(item) for item in value]
    else:
        text = clean_text(value)
        if not text:
            return []
        names = [item.strip() for item in re.split(r"[、,，;；\n]+", text) if item.strip()]

    return [name for name in names if name in image_map]


def render_images_in_context(context, tpl, image_files):
    image_map = {img.name: img for img in image_files} if image_files else {}

    for key, value in list(context.items()):
        image_names = split_image_names(value, image_map)
        if not image_names:
            if key in OPTIONAL_IMAGE_FIELDS:
                context[key] = f"{DELETE_PARAGRAPH_MARKER}{key}"
            continue

        images = []
        for name in image_names:
            image_map[name].seek(0)
            images.append(image_map[name])

        merged_img = make_a4_landscape_image(images)
        if merged_img is not None:
            context[key] = InlineImage(tpl, merged_img, width=Mm(LANDSCAPE_IMG_WIDTH_MM))

    return context


def remove_marked_paragraphs(doc, marker=DELETE_PARAGRAPH_MARKER):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    removed_count = 0
    for paragraph in paragraphs:
        if marker in paragraph.text:
            element = paragraph._element
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
                removed_count += 1

    return removed_count


def set_paragraph_text(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def has_device_table_placeholder(paragraph):
    return "{{环境试验用主要设备}}" in paragraph.text or DEVICE_TABLE_MARKER in paragraph.text


def fill_device_table(table, device_rows):
    table.style = "Table Grid"
    headers = ["设备名称", "出厂编号/版本", "证书编号", "有效期"]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row_data in device_rows:
        row = table.add_row().cells
        for idx, value in enumerate(row_data[:4]):
            row[idx].text = clean_text(value)


def iter_all_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph, None
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph, cell


def insert_device_table_after_marker(doc, device_rows, marker=DEVICE_TABLE_MARKER):
    if not device_rows:
        return 0

    inserted = 0
    for paragraph, cell in list(iter_all_paragraphs(doc)):
        if not has_device_table_placeholder(paragraph):
            continue
        if cell is not None:
            table = cell.add_table(rows=1, cols=4)
            fill_device_table(table, device_rows)
            parent = paragraph._element.getparent()
            if parent is not None:
                parent.remove(paragraph._element)
        else:
            table = doc.add_table(rows=1, cols=4)
            fill_device_table(table, device_rows)
            paragraph._p.addnext(table._tbl)
            parent = paragraph._element.getparent()
            if parent is not None:
                parent.remove(paragraph._element)
        inserted += 1
    return inserted


# ============================
# 上传文件
# ============================
st.markdown("""
<style>
.main .block-container {
    padding-top: 1.2rem;
    max-width: 1420px;
}
.hero-panel {
    padding: 30px 34px;
    border-radius: 20px;
    color: #182033;
    background:
        linear-gradient(120deg, rgba(42, 113, 196, 0.15), rgba(24, 169, 126, 0.12) 44%, rgba(255, 185, 86, 0.20)),
        linear-gradient(135deg, #f8fbff 0%, #eef9f3 54%, #fff7e8 100%);
    border: 1px solid rgba(76, 116, 154, 0.16);
    box-shadow: 0 18px 44px rgba(46, 79, 113, 0.14);
    margin-bottom: 16px;
    position: relative;
    overflow: hidden;
}
.hero-panel::after {
    content: "DOCX  XLSX  AI";
    position: absolute;
    right: 30px;
    top: 28px;
    color: rgba(45, 78, 108, 0.14);
    font-size: 2.2rem;
    font-weight: 800;
    letter-spacing: 0;
}
.hero-panel h1 {
    margin: 0;
    font-size: 2.25rem;
    letter-spacing: 0;
}
.hero-panel p {
    margin: 10px 0 0 0;
    color: #4a5872;
    font-size: 1.02rem;
}
.feature-grid {
    display: grid;
    grid-template-columns: repeat(4, minmax(0, 1fr));
    gap: 12px;
    margin: 12px 0 18px 0;
}
.feature-card {
    border-radius: 12px;
    padding: 14px 16px;
    min-height: 92px;
    border: 1px solid rgba(28, 59, 90, 0.10);
    box-shadow: 0 6px 18px rgba(38, 62, 87, 0.08);
}
.feature-card strong {
    display: block;
    font-size: 1rem;
    color: #1f2a3d;
    margin-bottom: 6px;
}
.feature-card span {
    color: #5a6679;
    font-size: 0.9rem;
}
.card-blue { background: #edf7ff; }
.card-green { background: #effaf2; }
.card-gold { background: #fff7e6; }
.card-rose { background: #fff0f3; }
.section-title {
    margin: 22px 0 12px 0;
    font-size: 1.1rem;
    color: #253348;
    font-weight: 700;
}
div[data-testid="stFileUploader"] section {
    border-radius: 14px;
    border: 1px dashed #b7c9df;
    background: #f7fbff;
    min-height: 92px;
}
div[data-testid="stExpander"] {
    border-radius: 14px;
    border: 1px solid #dce8f4;
    box-shadow: 0 8px 20px rgba(38, 62, 87, 0.06);
    background: #ffffff;
}
.stTabs [data-baseweb="tab-list"] {
    gap: 8px;
}
.stTabs [data-baseweb="tab"] {
    border-radius: 999px;
    padding: 8px 18px;
    background: #eef4fb;
}
.stTabs [aria-selected="true"] {
    background: #dff3eb;
    color: #0f654b;
}
div[data-testid="stDataFrame"], div[data-testid="stDataEditor"] {
    border-radius: 14px;
    overflow: hidden;
    border: 1px solid #e0e8f2;
}
.stButton > button, div[data-testid="stDownloadButton"] button {
    border-radius: 999px;
    border: 0;
    box-shadow: 0 6px 14px rgba(30, 82, 130, 0.12);
}
@media (max-width: 900px) {
    .feature-grid {
        grid-template-columns: repeat(2, minmax(0, 1fr));
    }
    .hero-panel::after {
        display: none;
    }
}
</style>
<div class="hero-panel">
  <h1>试验报告自动生成系统</h1>
  <p>上传委托单、模板和曲线图片，自动抽取字段、识别大纲与温湿度，并生成规范 Word 报告。</p>
</div>
<div class="feature-grid">
  <div class="feature-card card-blue"><strong>01 智能抽取</strong><span>从委托单自动带入产品、单位、试验项目等基础信息。</span></div>
  <div class="feature-card card-green"><strong>02 AI 填写</strong><span>从大纲、温湿度截图或采集网站提取试验条件和环境范围。</span></div>
  <div class="feature-card card-gold"><strong>03 曲线匹配</strong><span>识别真空、温度曲线和试验前后照片，自动写入字段。</span></div>
  <div class="feature-card card-rose"><strong>04 一键成稿</strong><span>按模板生成 Word 报告，自动清理空曲线位置。</span></div>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown("### AI 配置")
    global_api_key = st.text_input(
        "API Key",
        value=os.environ.get("OPENAI_API_KEY", ""),
        type="password",
        help="AI 识别功能使用；不会写入代码或报告。",
        key="openai_api_key",
    )
    st.caption(f"模型：{DEFAULT_AI_MODEL}")
    st.caption(f"接口：{DEFAULT_AI_BASE_URL}")

st.markdown('<div class="section-title">上传材料</div>', unsafe_allow_html=True)
upload_row_1 = st.columns(2)
with upload_row_1[0]:
    with st.container(border=True):
        st.markdown("#### 报告模板")
        st.caption("上传 Word 模板，程序会读取其中的占位字段。")
        tpl_file = st.file_uploader("Word 报告模板（.docx）", type=["docx"], label_visibility="collapsed")
with upload_row_1[1]:
    with st.container(border=True):
        st.markdown("#### 试验委托单")
        st.caption("上传委托单填写页，基础信息会自动带入。")
        data_file = st.file_uploader("试验委托单（.xlsx）", type=["xlsx"], label_visibility="collapsed")

upload_row_2 = st.columns(2)
with upload_row_2[0]:
    with st.container(border=True):
        st.markdown("#### 报告图片")
        st.caption("上传曲线图、真空图、试验前后照片，文件名用于生成报告。")
        img_files = st.file_uploader(
            "报告图片",
            type=["jpg", "jpeg", "png"],
            accept_multiple_files=True,
            label_visibility="collapsed",
        )
with upload_row_2[1]:
    with st.container(border=True):
        st.markdown("#### AI 识别附件")
        st.caption("上传试验大纲、温湿度截图、PDF/图片/Word/Excel/TXT。")
        ai_files = st.file_uploader(
            "AI 识别附件",
            type=["pdf", "jpg", "jpeg", "png", "webp", "docx", "xlsx", "xlsm", "txt", "md", "csv"],
            accept_multiple_files=True,
            label_visibility="collapsed",
        )
ai_image_files = [
    file
    for file in (ai_files or [])
    if Path(file.name).suffix.lower() in [".jpg", ".jpeg", ".png", ".webp"]
]
if ai_files:
    st.caption(f"AI 附件已上传 {len(ai_files)} 个；其中 {len(ai_image_files)} 张图片可用于温湿度识别。")


if tpl_file and data_file:
    try:
        tpl_file.seek(0)
        tpl = DocxTemplate(tpl_file)
        template_vars = sorted(tpl.get_undeclared_template_variables())

        context, key_order, commission_sheet_name, supplement_sheet_name = build_initial_context(data_file)
        for key in template_vars:
            context.setdefault(key, "")

        st.success(f"已从“{commission_sheet_name}”自动抽取委托单信息。")
        if supplement_sheet_name:
            st.caption(f"“{supplement_sheet_name}”仅作为补充来源，自动抽取不到的字段才会读取它。")

        data_signature = uploaded_file_signature(tpl_file, data_file)
        editor_df = ensure_editor_cache(context, template_vars, key_order, data_signature)
        apply_pending_widget_sync()

        extracted_fields = [
            *AUTO_FROM_COMMISSION_FIELDS,
            "产品尺寸",
            "技术状态",
            "生产单位",
            "送样方式",
            "贮存要求",
            "报告结论形式",
            "报告提交方式",
            "样品性能测试",
            "产品处置方式",
            "联系人",
            "联系电话",
            "试验条件",
            "特殊要求",
        ]
        extracted_preview = editor_df[editor_df["字段"].isin(extracted_fields)]
        with st.expander("从委托单自动带入的信息", expanded=False):
            st.dataframe(extracted_preview, hide_index=True, use_container_width=True)

        with st.expander("按填写要求补充信息", expanded=True):
            api_key = global_api_key
            config_cols = st.columns([2, 1, 1])
            config_cols[0].text_input(
                "AI 接口地址",
                value=DEFAULT_AI_BASE_URL,
                disabled=True,
                key="ai_base_url_display",
            )
            config_cols[1].text_input(
                "AI 模型",
                value=DEFAULT_AI_MODEL,
                disabled=True,
                key="ai_global_model_display",
            )
            config_cols[2].text_input(
                "推理强度",
                value=DEFAULT_AI_REASONING_EFFORT,
                disabled=True,
                key="ai_reasoning_effort_display",
            )
            tabs = st.tabs(["基本信息", "填写", "曲线名称"])

            with tabs[0]:
                st.subheader("编号")
                number_context = current_context_from_editor()
                current_base_number = normalize_base_report_number(
                    number_context.get("报告编号") or number_context.get("委托单号")
                )
                base_number = st.text_input(
                    "报告/委托基础编号",
                    value=current_base_number,
                    placeholder="例如：SYZX-ZKR-2026073",
                    help="填写基础编号后，报告编号自动加 c，委托单号自动加 b。",
                    key="linked_base_report_number",
                )
                update_editor_value("报告编号", linked_report_number(base_number), sync_widget=True)
                update_editor_value("委托单号", linked_order_number(base_number), sync_widget=True)
                if linked_experiment_number(base_number):
                    st.session_state["temperature_control_experiment_number"] = linked_experiment_number(base_number)
                st.caption(
                    f"报告编号：{linked_report_number(base_number) or '未填写'}；"
                    f"委托单号：{linked_order_number(base_number) or '未填写'}；"
                    f"控温说明编号：{linked_experiment_number(base_number) or '未填写'}"
                )

                st.subheader("日期")
                date_context = current_context_from_editor()
                range_cols = st.columns(2)
                for field in DATE_RANGE_FIELDS:
                    current_range = parse_date_range_for_input(date_context.get(field))
                    start_current = current_range[0] if current_range else None
                    end_current = current_range[1] if current_range else None
                    start_date = range_cols[0].date_input(
                        f"{field}开始",
                        value=start_current,
                        format="YYYY-MM-DD",
                        key=f"date_start_{field}",
                    )
                    end_date = range_cols[1].date_input(
                        f"{field}结束",
                        value=end_current,
                        format="YYYY-MM-DD",
                        key=f"date_end_{field}",
                    )
                    update_editor_value(field, format_date_value((start_date, end_date)))

                single_cols = st.columns(3)
                for index, field in enumerate(DATE_FIELDS):
                    current = parse_date_for_input(date_context.get(field))
                    selected_date = single_cols[index % 3].date_input(
                        field,
                        value=current,
                        format="YYYY-MM-DD",
                        key=f"date_{field}",
                    )
                    update_editor_value(field, format_date_value(selected_date))

                st.subheader("选择")
                option_context = current_context_from_editor()
                option_cols = st.columns(3)
                for index, (field, options) in enumerate(SELECT_FIELD_OPTIONS.items()):
                    current = clean_text(option_context.get(field))
                    option_list = list(options)
                    if current and current not in option_list:
                        option_list.insert(0, current)
                    selected = option_cols[index % 3].selectbox(
                        field,
                        options=option_list,
                        index=option_list.index(current) if current in option_list else 0,
                        key=f"select_{field}",
                    )
                    update_editor_value(field, selected)

                st.subheader("设备选择")
                device_context = current_context_from_editor()
                ledger_device_options = device_ledger_options()
                device_options = ledger_device_options or list(DEVICE_VERSION_OPTIONS.keys())
                current_device = device_options[0] if device_options else "KM3L"
                flow_name = clean_text(device_context.get("流程控制软件"))
                selected_main_device = clean_text(device_context.get("主试验设备") or device_context.get("环境试验用主要设备"))
                if selected_main_device in device_options:
                    current_device = selected_main_device
                for device in device_options:
                    if flow_name.startswith(device):
                        current_device = device
                        break

                selected_devices = st.multiselect(
                    "设备",
                    options=device_options,
                    default=[current_device] if current_device in device_options else device_options[:1],
                    key="multi_device_version",
                )
                if not selected_devices and current_device in device_options:
                    selected_devices = [current_device]
                if selected_devices:
                    update_editor_value("主试验设备", "、".join(selected_devices))
                    update_editor_value("环境试验用主要设备", "、".join(selected_devices))
                else:
                    update_editor_value("主试验设备", "")
                    update_editor_value("环境试验用主要设备", "")

                version_rows = []
                for device in selected_devices:
                    version_key = DEVICE_VERSION_ALIAS.get(device, device)
                    version_items = DEVICE_VERSION_OPTIONS.get(version_key, {})
                    if version_items:
                        for field, value in version_items.items():
                            update_editor_value(field, value)
                        version_rows.extend([{"字段": field, "内容": value} for field, value in version_items.items()])

                if version_rows:
                    st.dataframe(pd.DataFrame(version_rows), hide_index=True, use_container_width=True)

                device_selection_df = build_device_selection_dataframe(selected_devices)
                edited_device_selection_df = st.data_editor(
                    device_selection_df,
                    hide_index=True,
                    use_container_width=True,
                    column_config={
                        "插入": st.column_config.CheckboxColumn("插入", default=False),
                        "设备名称": st.column_config.TextColumn("设备名称", disabled=True),
                        "出厂编号/版本": st.column_config.TextColumn("出厂编号/版本", disabled=True),
                        "证书编号": st.column_config.TextColumn("证书编号", disabled=True),
                        "有效期": st.column_config.TextColumn("有效期", disabled=True),
                    },
                    key="device_selection_editor",
                )
                st.session_state["selected_device_table_rows"] = edited_device_selection_df[
                    edited_device_selection_df["插入"].astype(bool)
                ][["设备名称", "出厂编号/版本", "证书编号", "有效期"]].values.tolist()
                if not edited_device_selection_df.empty:
                    st.caption("环境试验用主要设备表（勾选后插入报告）")
                else:
                    st.warning("未在设备计量台账中找到所选设备的计量记录，请检查台账或设备名称。")

                st.subheader("人员")
                people_context = current_context_from_editor()
                leader_current = clean_text(people_context.get("试验技术负责人"))
                leader_options = list(PERSONNEL_OPTIONS)
                if leader_current and leader_current not in leader_options:
                    leader_options.insert(0, leader_current)
                leader = st.selectbox(
                    "试验技术负责人",
                    options=leader_options,
                    index=leader_options.index(leader_current) if leader_current in leader_options else 0,
                    key="select_试验技术负责人",
                )
                update_editor_value("试验技术负责人", leader)

                participant_current = split_people(people_context.get("试验参加者"))
                participant_options = list(PERSONNEL_OPTIONS)
                for name in participant_current:
                    if name not in participant_options:
                        participant_options.insert(0, name)
                participants = st.multiselect(
                    "试验参加者",
                    options=participant_options,
                    default=[name for name in participant_current if name in participant_options],
                    key="multi_试验参加者",
                )
                update_editor_value("试验参加者", "、".join(participants))

            with tabs[1]:
                st.subheader("AI 提取填写信息")
                outline_context = current_context_from_editor()
                target_test_item = clean_text(outline_context.get("试验项目"))
                target_test_name = clean_text(outline_context.get("试验名称"))
                target_experiment_number = find_experiment_number(outline_context)
                linked_control_number = linked_experiment_number(
                    st.session_state.get("linked_base_report_number")
                    or outline_context.get("报告编号")
                    or outline_context.get("委托单号")
                )
                default_experiment_number = linked_control_number or target_experiment_number
                if default_experiment_number and not clean_text(st.session_state.get("temperature_control_experiment_number")):
                    st.session_state["temperature_control_experiment_number"] = default_experiment_number
                manual_experiment_number = st.text_input(
                    "控温说明试验编号",
                    value=default_experiment_number,
                    placeholder="例如：SYZX-ZKR-2026-073",
                    help="默认由基本信息中的基础编号自动转换：SYZX-ZKR-2026073 -> SYZX-ZKR-2026-073；也可手动修改。",
                    key="temperature_control_experiment_number",
                )
                active_experiment_number = clean_text(manual_experiment_number) or default_experiment_number
                st.caption(
                    "检索条件："
                    + f"试验项目={target_test_item or '未填写'}；"
                    + f"试验名称={target_test_name or '未填写'}；"
                    + f"试验编号={active_experiment_number or '未识别'}"
                )
                if not target_test_item and not target_test_name:
                    st.warning("当前委托单未识别到试验项目/试验名称，AI 将无法可靠筛选大纲中的对应试验。")
                if not active_experiment_number:
                    st.info("未填写或识别到 SYZX-... 格式试验编号，将跳过控温说明网站，使用上传的大纲附件提取。")
                selected_test_hall = clean_text(outline_context.get("试验地点"))
                hall_for_temp = selected_test_hall if selected_test_hall in ["一区试验大厅", "二区试验大厅"] else "二区试验大厅"
                st.caption(f"温湿度将按“试验地点”：{hall_for_temp} 写入。")
                temp_start_date, temp_end_date = temp_humidity_date_range_from_context(outline_context)
                with st.container(border=True):
                    st.markdown("#### 温湿度获取")
                    temp_source = st.radio(
                        "温湿度来源",
                        options=["不处理温湿度", "网站获取", "截图识别"],
                        horizontal=True,
                        key="humidity_source",
                    )
                    site_cols = st.columns([1, 1, 1, 1])
                    site_cols[0].date_input(
                        "开始日期",
                        value=temp_start_date,
                        disabled=True,
                        help="自动使用基本信息里的试验准备时间；没有时使用试验日期起始日。",
                        key="humidity_site_start_date_display",
                    )
                    site_cols[1].date_input(
                        "结束日期",
                        value=temp_end_date,
                        disabled=True,
                        help="自动使用基本信息里的撤收时间；没有时使用试验日期结束日。",
                        key="humidity_site_end_date_display",
                    )
                    temp_site_org = site_cols[2].text_input(
                        "组织（可选）",
                        value=os.environ.get("TEMP_HUMIDITY_ORG_NAME", "环境试验中心"),
                        help="需与网页登录后的组织一致；你当前网页显示为环境试验中心，可保持默认。",
                        key="humidity_site_org",
                    )
                    site_cols[3].text_input(
                        "接口地址",
                        value=TEMP_HUMIDITY_DATA_URL,
                        disabled=True,
                        key="humidity_site_url_display",
                    )
                    credential_cols = st.columns([1, 1, 1])
                    temp_site_username = credential_cols[0].text_input(
                        "网站登录名",
                        value=os.environ.get("TEMP_HUMIDITY_USERNAME", ""),
                        help="填写登录页的用户名，一般是邮箱前缀/登录名，不一定是中文姓名。",
                        key="humidity_site_username",
                    )
                    temp_site_password = credential_cols[1].text_input(
                        "网站密码",
                        value=os.environ.get("TEMP_HUMIDITY_PASSWORD", ""),
                        type="password",
                        key="humidity_site_password",
                    )
                    temp_site_token = credential_cols[2].text_input(
                        "Basic Token（可选）",
                        value=os.environ.get("TEMP_HUMIDITY_BASIC_TOKEN", ""),
                        type="password",
                        help="已有 token 时可不填账号密码。",
                        key="humidity_site_token",
                    )
                    if temp_source == "截图识别":
                        st.caption("截图识别会使用“AI 识别附件”中的温湿度截图。")
                    elif temp_source == "网站获取":
                        st.caption("网站获取依赖主系统登录接口；若主系统连接超时，可先选“不处理温湿度”，只提取控温说明字段。")
                    if st.button("仅获取温湿度并写入", type="secondary", key="fetch_humidity_site_only"):
                        try:
                            with st.spinner("正在从温湿度网站获取数据..."):
                                humidity_result = fetch_temperature_humidity_stats(
                                    str(temp_start_date),
                                    str(temp_end_date),
                                    token=temp_site_token,
                                    username=temp_site_username,
                                    password=temp_site_password,
                                    org_name=temp_site_org,
                                )
                                st.session_state["humidity_site_result"] = humidity_result
                                temp_text, humidity_text = format_humidity_for_hall(humidity_result, hall_for_temp)
                                if temp_text:
                                    update_editor_value("试验地点温度", temp_text, sync_widget=True)
                                if humidity_text:
                                    update_editor_value("试验地点湿度", humidity_text, sync_widget=True)
                            st.success(f"已按 {hall_for_temp} 写入网站温湿度。")
                        except Exception as exc:
                            st.error(str(exc))
                ai_cols = st.columns([1, 1, 2])
                ai_model = DEFAULT_AI_MODEL
                ai_cols[0].text_input(
                    "AI 模型（PRO）",
                    value=ai_model,
                    disabled=True,
                    key="ai_model_display",
                )
                ai_max_tokens = ai_cols[1].number_input(
                    "输出 token 上限",
                    min_value=300,
                    max_value=4000,
                    value=1800,
                    step=100,
                    key="ai_max_tokens",
                )
                if ai_cols[2].button("AI 提取并写入", type="secondary"):
                    try:
                        wrote_fields = []
                        st.session_state.pop("temperature_control_error", None)
                        st.session_state.pop("humidity_fetch_error", None)
                        with st.spinner("正在识别试验大纲和温湿度..."):
                            experiment_number = clean_text(st.session_state.get("temperature_control_experiment_number")) or find_experiment_number(outline_context)
                            control_instruction = {}
                            if experiment_number:
                                try:
                                    control_instruction = fetch_temperature_control_instruction(experiment_number)
                                    if control_instruction:
                                        st.session_state["temperature_control_instruction"] = control_instruction
                                except Exception as exc:
                                    st.session_state["temperature_control_error"] = str(exc)
                            control_result = {}
                            outline_result = {}
                            if control_instruction:
                                control_result = extract_temperature_control_info_with_ai(
                                    control_instruction,
                                    api_key,
                                    requirement_context=outline_context,
                                    model=ai_model,
                                    max_output_tokens=int(ai_max_tokens),
                                )
                            if ai_files:
                                outline_result = extract_outline_info_with_ai(
                                    ai_files,
                                    api_key,
                                    requirement_context=outline_context,
                                    model=ai_model,
                                    max_output_tokens=int(ai_max_tokens),
                                )
                                if not outline_result.get("来源"):
                                    outline_result["来源"] = "上传大纲"
                            if control_result or outline_result:
                                merged_result = merge_ai_results(control_result, outline_result)
                                st.session_state["ai_outline_result"] = merged_result
                                for field in AI_EXTRACT_FIELDS:
                                    value = merged_result.get(field, "")
                                    if has_value(value):
                                        update_editor_value(field, value, sync_widget=True)
                                        wrote_fields.append(field)
                            elif experiment_number and st.session_state.get("temperature_control_error"):
                                st.session_state["temperature_control_error"] = (
                                    st.session_state["temperature_control_error"]
                                    + " 未上传大纲附件，无法回退提取。"
                                )
                            if temp_source == "网站获取":
                                try:
                                    humidity_result = fetch_temperature_humidity_stats(
                                        str(temp_start_date),
                                        str(temp_end_date),
                                        token=temp_site_token,
                                        username=temp_site_username,
                                        password=temp_site_password,
                                        org_name=temp_site_org,
                                    )
                                    st.session_state["humidity_site_result"] = humidity_result
                                    temp_text, humidity_text = format_humidity_for_hall(humidity_result, hall_for_temp)
                                    if temp_text:
                                        update_editor_value("试验地点温度", temp_text, sync_widget=True)
                                        wrote_fields.append("试验地点温度")
                                    if humidity_text:
                                        update_editor_value("试验地点湿度", humidity_text, sync_widget=True)
                                        wrote_fields.append("试验地点湿度")
                                except Exception as exc:
                                    st.session_state["humidity_fetch_error"] = str(exc)
                            elif temp_source == "截图识别" and ai_image_files:
                                try:
                                    humidity_result = extract_humidity_info_with_ai(
                                        ai_image_files,
                                        api_key,
                                        model=ai_model,
                                        max_output_tokens=800,
                                    )
                                    st.session_state["humidity_ai_result"] = humidity_result
                                    temp_text, humidity_text = format_humidity_for_hall(humidity_result, hall_for_temp)
                                    if temp_text:
                                        update_editor_value("试验地点温度", temp_text, sync_widget=True)
                                        wrote_fields.append("试验地点温度")
                                    if humidity_text:
                                        update_editor_value("试验地点湿度", humidity_text, sync_widget=True)
                                        wrote_fields.append("试验地点湿度")
                                except Exception as exc:
                                    st.session_state["humidity_fetch_error"] = str(exc)
                        if wrote_fields:
                            st.success("AI 提取完成，已写入：" + "、".join(wrote_fields))
                        else:
                            st.warning("AI 未提取到可写入内容，请检查上传附件或识别结果。")
                    except Exception as exc:
                        st.error(str(exc))

                ai_result = st.session_state.get("ai_outline_result")
                if ai_result:
                    source_name = clean_text(ai_result.get("来源"))
                    source_number = clean_text(ai_result.get("试验编号"))
                    if source_name:
                        st.success(f"AI 提取来源：{source_name}" + (f"（{source_number}）" if source_number else ""))
                    matched_item = clean_text(ai_result.get("匹配到的试验项目"))
                    if matched_item:
                        st.info(f"AI 匹配到的大纲试验项目：{matched_item}")
                    preview_rows = [
                        {"字段": field, "AI提取内容": ai_result.get(field, "")}
                        for field in AI_EXTRACT_FIELDS
                    ]
                    st.dataframe(pd.DataFrame(preview_rows), hide_index=True, use_container_width=True)
                    confirm_items = ai_result.get("需要人工确认", [])
                    if confirm_items:
                        st.warning("需要人工确认：" + "；".join(map(str, confirm_items)))
                control_error = clean_text(st.session_state.get("temperature_control_error"))
                if control_error:
                    st.warning(f"控温说明网站查询失败，已改用上传附件：{control_error}")
                humidity_error = clean_text(st.session_state.get("humidity_fetch_error"))
                if humidity_error:
                    st.warning(f"温湿度获取失败，已跳过温湿度写入，不影响其它 AI 字段：{humidity_error}")
                humidity_result = st.session_state.get("humidity_ai_result")
                if humidity_result:
                    preview = []
                    for hall_name in ["一区试验大厅", "二区试验大厅"]:
                        temp_text, humidity_text = format_humidity_for_hall(humidity_result, hall_name)
                        preview.append({
                            "大厅": hall_name,
                            "温度范围": temp_text,
                            "湿度范围": humidity_text,
                        })
                    st.dataframe(pd.DataFrame(preview), hide_index=True, use_container_width=True)
                    confirm_items = humidity_result.get("需要人工确认", []) if isinstance(humidity_result, dict) else []
                    if confirm_items:
                        st.warning("温湿度需要人工确认：" + "；".join(map(str, confirm_items)))

                humidity_site_result = st.session_state.get("humidity_site_result")
                if humidity_site_result:
                    st.markdown("#### 网站温湿度结果")
                    preview = []
                    for hall_name in ["一区试验大厅", "二区试验大厅"]:
                        temp_text, humidity_text = format_humidity_for_hall(humidity_site_result, hall_name)
                        hall_data = humidity_site_result.get(hall_name, {}) if isinstance(humidity_site_result, dict) else {}
                        preview.append({
                            "大厅": hall_name,
                            "设备名称": hall_data.get("设备名称", ""),
                            "温度范围": format_humidity_preview_value(temp_text),
                            "湿度范围": format_humidity_preview_value(humidity_text),
                            "备注": hall_data.get("备注", ""),
                        })
                    st.dataframe(pd.DataFrame(preview), hide_index=True, use_container_width=True)
                    confirm_items = humidity_site_result.get("需要人工确认", []) if isinstance(humidity_site_result, dict) else []
                    if confirm_items:
                        st.warning("网站温湿度需要人工确认：" + "；".join(map(str, confirm_items)))

                text_context = current_context_from_editor()
                text_cols = st.columns(3)
                for index, field in enumerate(MANUAL_TEXT_FIELDS):
                    current = clean_text(text_context.get(field))
                    value = text_cols[index % 3].text_area(
                        field,
                        value=current,
                        height=90 if field in ["试验目的", "试验过程"] else 68,
                        key=f"text_{field}",
                    )
                    update_editor_value(field, value)

            with tabs[2]:
                st.subheader("AI 识别曲线与照片名称")
                image_ai_cols = st.columns([1, 1, 2])
                image_ai_model = DEFAULT_AI_MODEL
                image_ai_cols[0].text_input(
                    "AI 模型（PRO）",
                    value=image_ai_model,
                    disabled=True,
                    key="image_ai_model_display",
                )
                image_ai_tokens = image_ai_cols[1].number_input(
                    "输出 token 上限",
                    min_value=500,
                    max_value=4000,
                    value=1800,
                    step=100,
                    key="image_ai_tokens",
                )
                if image_ai_cols[2].button("AI 识别并填写图片名称", type="secondary"):
                    try:
                        with st.spinner("正在识别图片类型并匹配字段..."):
                            st.session_state["image_ai_result"] = analyze_report_images_with_ai(
                                img_files,
                                current_context_from_editor(),
                                api_key,
                                model=image_ai_model,
                                max_output_tokens=int(image_ai_tokens),
                            )
                        st.success("图片识别完成，请检查下方文件名匹配结果后写入。")
                    except Exception as exc:
                        st.error(str(exc))

                image_ai_result = st.session_state.get("image_ai_result")
                if image_ai_result:
                    match_rows = [
                        {"字段": field, "匹配图片": "、".join(names)}
                        for field, names in image_ai_result.get("图片匹配", {}).items()
                        if names
                    ]
                    if match_rows:
                        st.dataframe(pd.DataFrame(match_rows), hide_index=True, use_container_width=True)

                    satisfaction = image_ai_result.get("是否满足要求", {})
                    if satisfaction:
                        st.dataframe(
                            pd.DataFrame([
                                {"字段": field, "AI判断": value}
                                for field, value in satisfaction.items()
                            ]),
                            hide_index=True,
                            use_container_width=True,
                        )

                    confirm_items = image_ai_result.get("需要人工确认", [])
                    if confirm_items:
                        st.warning("需要人工确认：" + "；".join(map(str, confirm_items)))

                    if st.button("将图片名称写入对应字段", type="primary"):
                        updates = {}
                        for field, names in image_ai_result.get("图片匹配", {}).items():
                            if names:
                                if field.startswith("试验温度曲线"):
                                    updates[field] = names[0]
                                else:
                                    updates[field] = "、".join(names)
                        for field, value in image_ai_result.get("是否满足要求", {}).items():
                            value = normalize_satisfaction_value(value)
                            if value:
                                updates[field] = value
                        if updates:
                            queue_widget_sync(updates)
                            st.success("已写入图片名称，页面将刷新显示。")
                            rerun_app()
                        else:
                            st.warning("没有可写入的识别结果。")

                name_context = current_context_from_editor()
                name_cols = st.columns(4)
                for index, field in enumerate(NAME_TEXT_FIELDS):
                    current = clean_text(name_context.get(field))
                    value = name_cols[index % 4].text_input(
                        field,
                        value=current,
                        key=f"name_{field}",
                    )
                    update_editor_value(field, value)

        editor_df = st.session_state["editor_df"]
        edited_df = st.data_editor(
            editor_df,
            hide_index=True,
            use_container_width=True,
            num_rows="fixed",
            disabled=["序号", "字段"],
            column_config={
                "序号": st.column_config.NumberColumn("序号", width="small"),
                "字段": st.column_config.TextColumn("字段", width="medium"),
                "内容": st.column_config.TextColumn("内容", width="large"),
            },
            key="context_editor",
        )
        st.session_state["editor_df"] = edited_df

        edited_context = editor_rows_to_context(edited_df)
        missing_keys = [
            key
            for key in template_vars
            if not has_value(edited_context.get(key))
            and not key.startswith("试验温度曲线")
            and key != "真空度曲线"
        ]
        if missing_keys:
            st.warning("以下模板字段还没有内容：" + "、".join(missing_keys))

        if st.button("生成报告", type="primary"):
            tpl_file.seek(0)
            tpl = DocxTemplate(tpl_file)
            selected_device_text = clean_text(edited_context.get("主试验设备") or edited_context.get("环境试验用主要设备"))
            selected_devices = [item.strip() for item in re.split(r"[、,，;；\n]+", selected_device_text) if item.strip()]
            device_rows = st.session_state.get("selected_device_table_rows", [])
            if device_rows:
                edited_context["环境试验用主要设备"] = DEVICE_TABLE_MARKER
            else:
                edited_context["环境试验用主要设备"] = ""
            render_context = render_images_in_context(edited_context, tpl, img_files)

            tpl.render(render_context)

            out_buf = BytesIO()
            tpl.save(out_buf)
            out_buf.seek(0)

            doc = Document(out_buf)
            removed_positions = remove_marked_paragraphs(doc)
            inserted_device_tables = insert_device_table_after_marker(doc, device_rows)
            out_buf = BytesIO()
            doc.save(out_buf)
            out_buf.seek(0)

            if removed_positions:
                st.success(f"报告生成成功，已删除 {removed_positions} 个空曲线位置。")
            else:
                st.success("报告生成成功。")
            if inserted_device_tables:
                st.info(f"已插入环境试验用主要设备表 {inserted_device_tables} 处。")
            st.download_button(
                "下载报告",
                data=out_buf,
                file_name=f"试验报告_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    except Exception as e:
        st.error(f"生成失败：{e}")
